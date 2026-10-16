"""Extract private supplier reference inputs without changing or publishing source workbooks.

Requires openpyxl and xlrd for read-only extraction. The output contains business
data and must remain in an ignored private directory (normally .data).
"""
from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

SCHEMAS = {
    "national": ["RFC", "Operaciones virtuales", "Estatus"],
    "foreign": ["Nombre o Razón social", "Apellido Paterno", "Apellido Materno", "País ",
                "Domicilio", "Cliente/Proveedor", "Estatus"],
}
TABLE_CODES = ("501", "505", "507")


def file_source(path: Path) -> dict:
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    return {"file": path.name, "sha256": digest}


def read_template(path: Path, kind: str) -> tuple[list[str], list[dict]]:
    """Keep every original supplier, every column, order, and status verbatim."""
    xlrd = importlib.import_module("xlrd")
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = book.sheet_by_index(0)
        headers = [str(cell) for cell in sheet.row_values(0)]
        if headers != SCHEMAS[kind]:
            raise ValueError(f"Unexpected {kind} template headers: {headers!r}")
        source = file_source(path)
        rows = []
        for row_number in range(1, sheet.nrows):
            values = sheet.row_values(row_number)
            if not any(str(value).strip() for value in values):
                continue
            record = dict(zip(headers, values, strict=True))
            record.update(status=record["Estatus"], originalStatus=record["Estatus"],
                          source=source | {"sheet": sheet.name, "row": row_number + 1})
            if kind == "national":
                value = str(record["Operaciones virtuales"]).strip().casefold()
                if value not in {"si", "sí", "no"}:
                    raise ValueError(f"Unrecognized virtual-operation confirmation in {path.name}:{row_number + 1}")
                record.update(rfc=str(record["RFC"]).strip(), virtualOperations=value in {"si", "sí"})
            else:
                record.update(name=record[headers[0]], paternalSurname=record[headers[1]],
                              maternalSurname=record[headers[2]], country=record[headers[3]],
                              address=record[headers[4]], role=record[headers[5]])
            rows.append(record)
        return headers, rows
    finally:
        book.release_resources()


def json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def read_annual(path: Path) -> dict:
    """Read the regenerated annual's traceable rows; never create monthly runs."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    source = file_source(path)
    tables = {code: [] for code in TABLE_CODES}
    try:
        for sheet in book:
            code = sheet.title.split()[0]
            if code not in tables:
                continue
            iterator = sheet.iter_rows(values_only=True)
            headers = list(next(iterator, ()))
            if not {"Periodo", "ArchivoOrigen", "FolioOrigen", "Patente", "Pedimento", "SeccionAduanera"}.issubset(headers):
                raise ValueError(f"Missing trace or declaration columns in {sheet.title}")
            for row_number, values in enumerate(iterator, 2):
                if not any(value is not None for value in values):
                    continue
                record = {str(header): json_value(value) for header, value in zip(headers, values, strict=True)}
                record["source"] = source | {"sheet": sheet.title, "row": row_number}
                tables[code].append(record)
        if any(not rows for rows in tables.values()):
            raise ValueError("The baseline requires populated 501, 505, and 507 tables")
        return {"source": source, "tables": tables}
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--national", required=True, type=Path)
    parser.add_argument("--foreign", required=True, type=Path)
    parser.add_argument("--annual", required=True, type=Path)
    parser.add_argument("--scope", required=True)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--xls-reader-path", type=Path, help="Optional isolated xlrd module directory")
    args = parser.parse_args()
    if args.xls_reader_path:
        sys.path.insert(0, str(args.xls_reader_path.resolve()))
    schemas, records = {}, {}
    for kind in SCHEMAS:
        schemas[kind], records[kind] = read_template(getattr(args, kind), kind)
    payload = {"schemaVersion": 1, "scopeId": args.scope, "schemas": schemas, **records,
               "baseline": read_annual(args.annual)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(args.output.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    temporary.replace(args.output)
    print(json.dumps({"rows": {kind: len(rows) for kind, rows in records.items()},
                      "statuses": {kind: dict(Counter(row["status"] for row in rows)) for kind, rows in records.items()},
                      "baselineRows": {code: len(rows) for code, rows in payload["baseline"]["tables"].items()}},
                     ensure_ascii=True))


if __name__ == "__main__":
    main()
