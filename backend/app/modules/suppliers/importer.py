"""Read the two user-selected templates without changing the source files."""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from app.core.errors import ApplicationError
from .service import DEFAULT_SCHEMAS, normalized

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def read_list(content: bytes, filename: str, kind: str):
    return read_workbook(content, filename, kind)[0]


def read_workbook(content: bytes, filename: str, kind: str):
    from .workbooks import encode_template, from_xls
    filename = Path(filename.replace("\\", "/")).name
    if not content or len(content) > MAX_BYTES:
        raise ApplicationError(400, "SUPPLIER_IMPORT_SIZE", "Cada lista debe contener datos y pesar hasta 5 MB")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True)
            try:
                sheet = book.sheet_by_index(0)
                if sheet.nrows > MAX_ROWS + 1:
                    raise ValueError("row limit")
                title, rows = sheet.name, [sheet.row_values(i, 0, len(DEFAULT_SCHEMAS[kind]) + 1) for i in range(sheet.nrows)]
                template = encode_template(from_xls(book))
            finally:
                book.release_resources()
        elif extension == ".xlsx":
            # Bound expanded input before the XML reader touches the workbook.
            with ZipFile(BytesIO(content)) as zipped:
                if len(zipped.infolist()) > 256 or sum(i.file_size for i in zipped.infolist()) > 4 * MAX_BYTES:
                    raise ValueError("expanded size")
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(content), keep_links=False)
            try:
                sheet = book.worksheets[0]
                if sheet.max_row is None or sheet.max_row > MAX_ROWS + 1:
                    raise ValueError("row limit")
                title, rows = sheet.title, []
                for row in sheet.iter_rows(max_row=sheet.max_row, max_col=len(DEFAULT_SCHEMAS[kind]) + 1, values_only=True):
                    if len(rows) > MAX_ROWS and any(value is not None for value in row):
                        raise ValueError("row limit")
                    rows.append(list(row))
                template = encode_template(book)
            finally:
                book.close()
        else:
            raise ValueError("unsupported extension")
    except ImportError as error:
        raise ApplicationError(503, "SUPPLIER_READER_UNAVAILABLE", "Falta actualizar el lector de Excel del backend. Revisa las dependencias xlrd y openpyxl") from error
    except Exception as error:
        raise ApplicationError(400, "INVALID_SUPPLIER_WORKBOOK", "No se pudo leer la lista. Selecciona un Excel .xls o .xlsx válido de hasta 5000 proveedores") from error
    headers = [str(value or "") for value in rows[0]] if rows else []
    # Excel may retain formatted trailing blank columns after saving as XLSX.
    while headers and not headers[-1]:
        headers.pop()
    if headers != DEFAULT_SCHEMAS[kind]:
        raise ApplicationError(400, "INVALID_SUPPLIER_SCHEMA", "Las columnas de las listas no coinciden con las listas originales")
    records = []
    digest = sha256(content).hexdigest()
    for index, values in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        record = {header: values[i] if i < len(values) and values[i] is not None else "" for i, header in enumerate(headers)}
        if kind == "national" and normalized(record["Operaciones virtuales"]) not in {"", "SI", "NO"}:
            raise ApplicationError(400, "INVALID_VIRTUAL_CONFIRMATION", f"Fila {index}: indica Si o No en Operaciones virtuales")
        if kind == "foreign" and normalized(record["Cliente/Proveedor"]) not in {"", "PROVEEDOR"}:
            raise ApplicationError(400, "INVALID_FOREIGN_SUPPLIER", f"Fila {index}: selecciona una lista de proveedores")
        record["source"] = {"file": filename, "sha256": digest, "sheet": title, "row": index}
        records.append(record)
    return records, template
