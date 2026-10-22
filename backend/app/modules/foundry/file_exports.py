"""Deterministic multi-format files from authorized queries and attachments."""
import csv
from datetime import datetime
from decimal import Decimal
from html import escape
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
import reportlab
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PDFImage, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from sqlalchemy import select
import xlsxwriter

from app.core.errors import ApplicationError
from app.persistence.models import AgentArtifact, StoredDocument
from .artifacts import FORMATS, MAX_ARTIFACT_BYTES, store_artifact
from .conversations import own_conversation
from .data_queries import MAX_EXPORT_CELLS, fail, query_data
from .file_visuals import display, draw_chart
from .presentation_export import render_presentation
from .tools import CreateFile, CreatePackage
from .uploads import query_file


def json_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def caption(result, language):
    parts = []
    for source in result.sources:
        if "attachmentId" in source:
            parts.append(f'{source["name"]}. {source["sheet"]}')
        else:
            version = "Version" if language == "en" else "Versión"
            parts.append(f'{source["period"]}. {version} {source["version"]}')
    if len(parts) > 3:
        parts = parts[:3] + [f"{len(parts) - 3} more sources" if language == "en" else f"{len(parts) - 3} fuentes más"]
    return "; ".join(parts)


def source_key(source):
    return source.get("runId") or source["attachmentId"] + ":" + source["sheet"]


def source_manifest(title, format, sections, language):
    sources = {}
    entries = []
    for section, result, _ in sections:
        tables = [result.specification["tableCode"], *[join["tableCode"] for join in result.specification.get("joins", [])]] if "tableCode" in result.specification else []
        for source in result.sources:
            key = source_key(source)
            merged = source | {"tables": sorted(set(sources.get(key, {}).get("tables", [])) | set(source.get("tables", [])) | set(tables))}
            sources[key] = merged
        entries.append({"title": section["title"], "headers": result.headers, "types": result.types,
            "rows": len(result.rows), "matchedSourceRows": result.matched_rows, "scannedSourceRows": result.scanned_rows,
            "query": result.specification, "diagnostics": result.diagnostics, "units": dict(zip(result.headers, result.units)),
            "sources": result.sources, "commentary": section["commentary"], "chart": section["chart"]})
    return {"title": title, "format": format, "language": language, "complete": True,
            "sections": entries, "sources": list(sources.values())}


def render_csv(result):
    def safe(value):
        if value is None:
            return ""
        raw = value.isoformat() if isinstance(value, datetime) else str(value)
        if isinstance(value, str) and raw.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@")):
            return "'" + raw
        return raw
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(result.headers)
    writer.writerows([[safe(value) for value in row] for row in result.rows])
    return stream.getvalue().encode("utf-8-sig")


def render_xlsx(sections, manifest, language):
    stream = BytesIO()
    preserved_as_text = 0
    with xlsxwriter.Workbook(stream, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}) as book:
        heading = book.add_format({"bold": True, "font_color": "white", "bg_color": "#145C56"})
        numbers = book.add_format({"num_format": "#,##0.##########"})
        dates = book.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
        for index, (section, result, _) in enumerate(sections, 1):
            name = re.sub(r"[\[\]:*?/\\']", "", section["title"])[:24] or "Data"
            page = book.add_worksheet(f"{index} {name}"[:31])
            for column, header in enumerate(result.headers):
                page.write_string(0, column, header, heading)
            for row_index, row in enumerate(result.rows, 1):
                for column, value in enumerate(row):
                    if value is None:
                        continue
                    if isinstance(value, datetime):
                        page.write_datetime(row_index, column, value, dates)
                    elif isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                        # Excel supports 15 significant digits; preserve source
                        # overprecision as literal text, following the base exporter.
                        from app.modules.engine.exporter import _number_cell
                        number = _number_cell(value)
                        if number is not None:
                            page.write_number(row_index, column, number, numbers)
                        else:
                            page.write_string(row_index, column, str(value))
                            preserved_as_text += 1
                    else:
                        page.write_string(row_index, column, str(value))
                        preserved_as_text += int(result.types[column] in {"number", "date"})
            page.freeze_panes(1, 0)
            page.autofilter(0, 0, len(result.rows), len(result.headers) - 1)
            page.set_column(0, len(result.headers) - 1, 24)
        page = book.add_worksheet("Sources" if language == "en" else "Fuentes")
        page.write_string(0, 0, "Section", heading)
        page.write_string(0, 1, "Sources, query and diagnostics", heading)
        for index, entry in enumerate(manifest["sections"], 1):
            serialized = json.dumps(entry, ensure_ascii=False)
            if len(serialized) > 32767:
                fail("agent_cell_too_large", "El manifiesto de una sección supera el tamaño de celda de Excel.")
            page.write_string(index, 0, entry["title"])
            page.write_string(index, 1, serialized)
        page.set_column(0, 0, 35)
        page.set_column(1, 1, 100)
        manifest["preservedAsTextCells"] = preserved_as_text
        page.write_string(len(sections) + 2, 0, f"Values preserved as text: {preserved_as_text}")
    return stream.getvalue()


def render_word(title, sections, manifest, language):
    document = Document()
    page = document.sections[0]
    if max(len(result.headers) for _, result, _ in sections) > 5:
        page.orientation = WD_ORIENT.LANDSCAPE
        page.page_width, page.page_height = page.page_height, page.page_width
    page.top_margin = page.bottom_margin = Inches(.7)
    page.left_margin = page.right_margin = Inches(.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Title"].font.size = Pt(30)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        for border in document.styles[name]._element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    document.styles["Heading 1"].font.color.rgb = RGBColor.from_string("145C56")
    document.add_heading(title, 0)
    for section, result, source_caption in sections:
        document.add_heading(section["title"], 1)
        if section["commentary"]:
            document.add_paragraph(section["commentary"])
        if section["chart"] != "none":
            picture = draw_chart(result, section["title"], section["chart"], "png", source_caption, manifest)
            document.add_picture(BytesIO(picture), width=page.page_width - page.left_margin - page.right_margin)
        table = document.add_table(rows=1, cols=len(result.headers))
        table.style = "Light Shading Accent 1"
        for cell, header in zip(table.rows[0].cells, result.headers):
            cell.text = header
            shade = OxmlElement("w:shd")
            shade.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "145C56")
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row in result.rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = display(value)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string("28414C")
        document.add_paragraph(("Source: " if language == "en" else "Fuente: ") + source_caption)
    document.add_page_break()
    document.add_heading("Sources and calculation details" if language == "en" else "Fuentes y detalle de cálculo", 1)
    for entry in manifest["sections"]:
        document.add_heading(entry["title"], 2)
        for source in entry["sources"]:
            if "attachmentId" in source:
                document.add_paragraph(f'{source["name"]}. {source["sheet"]}. SHA-256: {source["sha256"]}')
            else:
                document.add_paragraph(f'{source["period"]}. {"Version" if language == "en" else "Versión"} {source["version"]}. ID: {source["runId"]}')
        document.add_paragraph(("Result rows: " if language == "en" else "Filas del resultado: ") + str(entry["rows"]))
        document.add_paragraph(("Columns: " if language == "en" else "Columnas: ") + ", ".join(entry["headers"]))
        for item in entry["query"].get("filters", []):
            document.add_paragraph(f'{item["column"]}: {item["operator"]} {item["value"]}')
    document.core_properties.author = "DataStage"
    document.core_properties.title = title
    output = BytesIO()
    document.save(output)
    # Keep the full declarative plan in a standard custom XML part, alongside
    # the readable provenance appendix.
    from .presentation_export import REL, package_xml
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(BytesIO(output.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    relationships = ET.fromstring(parts["_rels/.rels"])
    ET.SubElement(relationships, "{" + REL + "}Relationship", {"Id": "rIdDataStageSources", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml", "Target": "customXml/datastage-sources.xml"})
    parts["_rels/.rels"] = package_xml(relationships, REL)
    parts["customXml/datastage-sources.xml"] = ('<sources xmlns="urn:datastage:sources">' + escape(json.dumps(manifest, ensure_ascii=False)) + '</sources>').encode()
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return output.getvalue()


def render_pdf(title, sections, manifest, language):
    output = BytesIO()
    styles = getSampleStyleSheet()
    # ReportLab ships these fonts. Embed them for portable ES/EN glyphs and
    # consistent character widths without requiring fonts on the server.
    for name, filename in (("DataStageSans", "Vera.ttf"), ("DataStageSansBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(Path(reportlab.__file__).parent / "fonts" / filename)))
    for style in styles.byName.values():
        if hasattr(style, "fontName"):
            style.fontName = "DataStageSansBold" if style.fontName.endswith("Bold") else "DataStageSans"
    for name in ("Heading1", "Heading2"):
        styles[name].textColor = colors.HexColor("#145C56")
    paper = (842, 595) if max(len(result.headers) for _, result, _ in sections) > 5 else (595, 842)
    doc = SimpleDocTemplate(output, pagesize=paper, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36,
                            title=title, author="DataStage")
    parts = [Paragraph(escape(title), styles["Title"]), Spacer(1, 16)]
    for section, result, source_caption in sections:
        parts.append(Paragraph(escape(section["title"]), styles["Heading1"]))
        if section["commentary"]:
            parts.append(Paragraph(escape(section["commentary"]).replace("\n", "<br/>"), styles["BodyText"]))
            parts.append(Spacer(1, 12))
        if section["chart"] != "none":
            data = draw_chart(result, section["title"], section["chart"], "png", source_caption, manifest)
            image_width, image_height = ImageReader(BytesIO(data)).getSize()
            width = paper[0] - 72
            height = width * image_height / image_width
            if height > paper[1] - 180:
                height = paper[1] - 180
                width = height * image_width / image_height
            parts.extend([PDFImage(BytesIO(data), width=width, height=height), Spacer(1, 12)])
        cells = [[Paragraph(escape(display(value)), styles["BodyText"]) for value in row]
                 for row in [result.headers, *result.rows]]
        table = LongTable(cells, repeatRows=1, colWidths=[(paper[0] - 72) / len(result.headers)] * len(result.headers))
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E1EEEB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F6")])]))
        parts.extend([table, Spacer(1, 10), Paragraph(escape(source_caption), styles["BodyText"]), Spacer(1, 18)])
    parts.append(Paragraph("Sources and calculation details" if language == "en" else "Fuentes y detalle de cálculo", styles["Heading1"]))
    for entry in manifest["sections"]:
        parts.append(Paragraph(escape(entry["title"]), styles["Heading2"]))
        for source in entry["sources"]:
            if "attachmentId" in source:
                description = f'{source["name"]}. {source["sheet"]}. SHA-256: {source["sha256"]}'
            else:
                description = f'{source["period"]}. {"Version" if language == "en" else "Versión"} {source["version"]}. ID: {source["runId"]}'
            parts.append(Paragraph(escape(description), styles["BodyText"]))
        parts.append(Paragraph(("Result rows: " if language == "en" else "Filas del resultado: ") + str(entry["rows"]), styles["BodyText"]))
        parts.append(Paragraph(escape(("Columns: " if language == "en" else "Columnas: ") + ", ".join(entry["headers"])), styles["BodyText"]))
        for item in entry["query"].get("filters", []):
            parts.append(Paragraph(escape(f'{item["column"]}: {item["operator"]} {item["value"]}'), styles["BodyText"]))
    doc.build(parts)
    return output.getvalue()


def render_html(title, sections, manifest, language):
    content = [f'<!doctype html><html lang="{language}"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
        '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; img-src data:; base-uri \'none\'; form-action \'none\'">',
        f'<title>{escape(title)}</title><style>body{{font:16px Arial,sans-serif;color:#28414c;max-width:1200px;margin:40px auto;padding:0 24px}}h1,h2{{color:#145c56}}table{{border-collapse:collapse;width:100%;margin:24px 0}}th,td{{padding:10px;text-align:left;border-bottom:1px solid #dce6e7;white-space:pre-wrap}}th{{background:#e4efed}}.table{{overflow:auto}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px}}svg{{max-width:100%;height:auto}}section{{margin:48px 0}}</style><h1>{escape(title)}</h1><nav>']
    for index, (section, _, _) in enumerate(sections):
        content.append(f'<p><a href="#section-{index}">{escape(section["title"])}</a></p>')
    content.append('</nav>')
    for index, (section, result, source_caption) in enumerate(sections):
        content.append(f'<section id="section-{index}"><h2>{escape(section["title"])}</h2><p>{escape(section["commentary"])}</p>')
        if section["chart"] != "none":
            content.append(draw_chart(result, section["title"], section["chart"], "svg", source_caption, manifest).decode())
        content.append('<div class="table"><table><thead><tr>' + ''.join(f'<th>{escape(value)}</th>' for value in result.headers) + '</tr></thead><tbody>')
        for row in result.rows:
            content.append('<tr>' + ''.join(f'<td>{escape(display(value))}</td>' for value in row) + '</tr>')
        content.append(f'</tbody></table></div><p>{escape(source_caption)}</p></section>')
    content.append('<details><summary>Sources / Fuentes</summary><pre>' + escape(json.dumps(manifest, ensure_ascii=False, indent=2)) + '</pre></details></html>')
    return ''.join(content).encode("utf-8")


def create_file(session, storage, principal, conversation_id, specification, language):
    args = CreateFile.model_validate(specification).model_dump()
    own_conversation(session, principal, conversation_id)
    format = args["format"]
    if format in {"csv", "png", "svg"} and len(args["sections"]) != 1:
        fail("agent_file_sections", "CSV, PNG y SVG requieren una sola sección.")
    sections = []
    cells = 0
    for section in args["sections"]:
        query = section["query"]
        result = query_file(session, storage, principal, conversation_id, query) if "fileId" in query else query_data(session, principal.scope_id, query)
        cells += (len(result.rows) + 1) * len(result.headers)
        if cells > MAX_EXPORT_CELLS or any(len(str(value)) > 32767 for row in [result.headers, *result.rows] for value in row):
            fail("agent_export_too_large", "El archivo supera los límites de celdas o texto; reduce su selección.", 413)
        if format in {"docx", "pptx", "pdf"}:
            if section["chart"] != "none" and len(result.rows) > 12:
                fail("agent_document_chart", "Las gráficas dentro de documentos admiten hasta 12 categorías legibles. Agrupa los datos o genera una imagen independiente.")
            if len(result.headers) > 8 or len(result.rows) > 200 or cells > 2000:
                fail("agent_document_too_large", "El documento visual admite 8 columnas, 200 filas por sección y 2,000 celdas totales. Agrupa los datos o utiliza CSV/JSON/XLSX.", 413)
            for row in [result.headers, *result.rows]:
                if any(len(display(value)) > 240 for value in row):
                    fail("agent_visual_cell", "Una celda es demasiado larga para el documento visual. Selecciona campos más breves o genera Excel/CSV.")
        if any(ord(char) < 32 and char not in "\t\n\r" for text in [args["title"], section["title"], section["commentary"], *result.headers] for char in text):
            fail("agent_invalid_character", "El documento contiene caracteres de control no admitidos.")
        sections.append((section, result, caption(result, language)))
    manifest = source_manifest(args["title"], format, sections, language)
    if format == "csv":
        content = render_csv(sections[0][1])
        manifest["csvFormulaProtection"] = "apostrophe_prefix_for_formula_like_text"
    elif format == "json":
        payload = {"manifest": manifest, "sections": [{"title": section["title"], "headers": result.headers,
            "types": result.types, "rows": [[json_value(value) for value in row] for row in result.rows]} for section, result, _ in sections]}
        content = json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")
    elif format == "xlsx":
        content = render_xlsx(sections, manifest, language)
    elif format == "docx":
        content = render_word(args["title"], sections, manifest, language)
    elif format == "pptx":
        content = render_presentation(args["title"], sections, manifest, language)
    elif format == "pdf":
        content = render_pdf(args["title"], sections, manifest, language)
    elif format == "html":
        content = render_html(args["title"], sections, manifest, language)
    else:
        section, result, source_caption = sections[0]
        content = draw_chart(result, args["title"], section["chart"], format, source_caption, manifest)
    return store_artifact(session, storage, principal, conversation_id, content, "." + format, args["title"], manifest)


def create_package(session, storage, principal, conversation_id, specification):
    args = CreatePackage.model_validate(specification)
    own_conversation(session, principal, conversation_id)
    if len(set(args.documentIds)) != len(args.documentIds):
        fail("agent_package_duplicate", "Selecciona cada archivo una sola vez.")
    entries, contents, size = [], [], 0
    for index, document_id in enumerate(args.documentIds, 1):
        row = session.execute(select(StoredDocument, AgentArtifact).join(AgentArtifact, AgentArtifact.document_id == StoredDocument.id)
            .where(StoredDocument.id == document_id, StoredDocument.scope_id == principal.scope_id,
                AgentArtifact.conversation_id == conversation_id, StoredDocument.status == "available")).first()
        if row is None or row[0].kind not in {kind for kind, _ in FORMATS.values()} or row[0].kind == "agent_zip":
            fail("agent_package_source", "Selecciona descargas generadas de esta conversación; no se incluyen ZIP anidados.")
        document, artifact = row
        size += document.size_bytes
        if size > 40 * 1024 * 1024:
            fail("agent_package_large", "El paquete supera 40 MiB de archivos de origen.", 413)
        try:
            content = storage.read_bytes(document.storage_key)
        except OSError as error:
            raise ApplicationError(503, "agent_artifact_unavailable", "Uno de los archivos ya no está disponible.") from error
        import hashlib
        if hashlib.sha256(content).hexdigest() != document.sha256:
            fail("agent_artifact_changed", "Un archivo cambió; vuelve a generarlo antes de empaquetar.", 503)
        name = f"{index:02d}_{document.original_name}"
        contents.append((name, content))
        entries.append({"name": name, "documentId": document.id, "sha256": document.sha256,
                        "sizeBytes": len(content), "manifest": json.loads(artifact.manifest_json)})
    sources = {}
    for entry in entries:
        for source in entry["manifest"].get("sources", []):
            key = source_key(source)
            sources[key] = source | {"tables": sorted(set(sources.get(key, {}).get("tables", [])) | set(source.get("tables", [])))}
    manifest = {"title": args.title, "complete": True, "files": entries, "sources": list(sources.values())}
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in contents:
            archive.writestr(name, content)
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    if output.tell() > MAX_ARTIFACT_BYTES:
        fail("agent_package_large", "El ZIP supera 20 MiB comprimidos.", 413)
    result = store_artifact(session, storage, principal, conversation_id, output.getvalue(), ".zip", args.title, manifest)
    # The ZIP and persisted record keep the full manifests. Do not repeat every
    # nested query in the model's bounded tool context.
    full = result["manifest"]
    result["manifest"] = {**full, "files": [{key: item[key] for key in ("name", "documentId", "sha256", "sizeBytes")} for item in full["files"]]}
    return result
