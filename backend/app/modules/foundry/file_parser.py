"""Bounded, text/data-only attachment reader. Run in a short-lived child process.

No macros, formulas, links, scripts or embedded objects are executed. Office
archives are inspected without extracting paths onto the filesystem.
"""
import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import json
from pathlib import Path, PurePosixPath
import re
import stat
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

MAX_ROWS = 50_000
MAX_COLUMNS = 80
MAX_CELLS = 200_000
MAX_TEXT = 2_000_000
MAX_SHEETS = 20
MAX_CELL = 50_000
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
FORMATS = {".xlsx", ".csv", ".pdf", ".docx", ".pptx"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
IDENTIFIERS = {"id", "rfc", "patente", "pedimento", "pedimentocompleto", "fraccion", "fraccionarancelaria", "tariffcode",
               "np", "pn", "partnumber", "factura", "numerofactura", "invoicenumber", "codigopostal", "postalcode", "zipcode"}


def normalize(value):
    text = "".join(char for char in unicodedata.normalize("NFKD", str(value)) if not unicodedata.combining(char))
    return re.sub(r"\W+", "", text).casefold()


def office_archive(content, suffix):
    archive = zipfile.ZipFile(BytesIO(content))
    members = archive.infolist()
    names, size = set(), 0
    if len(members) > 2000:
        raise ValueError("El archivo Office contiene demasiados componentes.")
    for member in members:
        path = PurePosixPath(member.filename.replace("\\", "/"))
        size += member.file_size
        if (path.is_absolute() or ".." in path.parts or ":" in member.filename or
                member.filename in names or member.flag_bits & 1 or stat.S_ISLNK(member.external_attr >> 16)
                or member.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED} or member.file_size > 20 * 1024 * 1024
                or size > MAX_ARCHIVE_BYTES or member.file_size > max(1, member.compress_size) * 200):
            raise ValueError("La estructura comprimida del archivo Office no está permitida.")
        names.add(member.filename)
        if (any(token in member.filename.casefold() for token in ("vbaproject", "activex/")) or
                "embeddings/" in member.filename.casefold() and not member.filename.lower().endswith(".xlsx")):
            raise ValueError("Utiliza un archivo Office sin macros, controles ni objetos incrustados.")
        if member.filename.casefold().endswith((".xml", ".rels")):
            raw = archive.read(member)
            if b"\x00" in raw or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                raise ValueError("No se admiten declaraciones de entidades en archivos Office.")
    required = {".xlsx": "xl/workbook.xml", ".docx": "word/document.xml", ".pptx": "ppt/presentation.xml"}[suffix]
    if required not in names or "[Content_Types].xml" not in names:
        raise ValueError("La extensión no coincide con el contenido del archivo.")
    return archive


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (bool, int, float)):
        return value
    value = str(value)
    if len(value) > MAX_CELL or any(ord(char) < 32 and char not in "\t\n\r" for char in value):
        raise ValueError("Una celda contiene texto demasiado largo o caracteres no admitidos.")
    return value


def sheet(name, raw_rows):
    rows = iter(raw_rows)
    raw_headers = next(rows, [])
    if not raw_headers:
        return {"name": name, "headers": [], "types": [], "rows": []}
    if len(raw_headers) > MAX_COLUMNS:
        raise ValueError("Una hoja supera 80 columnas. Reduce su rango de datos.")
    headers, used = [], set()
    for index, value in enumerate(raw_headers, 1):
        label = str(value or f"Column {index}").strip() or f"Column {index}"
        if len(label) > 100:
            raise ValueError("Una cabecera supera 100 caracteres.")
        candidate, ordinal = label, 2
        while normalize(candidate) in used:
            candidate = f"{label[:90]} ({ordinal})"
            ordinal += 1
        headers.append(candidate)
        used.add(normalize(candidate))
    output = []
    for raw in rows:
        if not any(value is not None and value != "" for value in raw):
            continue
        if len(raw) > len(headers) and any(value not in (None, "") for value in raw[len(headers):]):
            raise ValueError("Hay filas con más campos que las cabeceras.")
        if len(output) >= MAX_ROWS or (len(output) + 1) * len(headers) > MAX_CELLS:
            raise ValueError("El archivo supera el límite de filas o celdas; no se ha leído parcialmente.")
        output.append([clean(raw[index]) if index < len(raw) else None for index in range(len(headers))])
    types = []
    for index, header in enumerate(headers):
        values = [row[index] for row in output if row[index] not in (None, "")]
        if normalize(header) in IDENTIFIERS or normalize(header).endswith("id"):
            kind = "text"
            for row in output:
                if row[index] is not None:
                    row[index] = str(row[index])
        elif values and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
            kind = "number"
        elif values and all(isinstance(value, str) and re.fullmatch(r"\d{4}-\d\d-\d\d(?:T[\d:.]+)?", value) for value in values):
            kind = "date"
        else:
            kind = "text"
        types.append(kind)
    return {"name": name, "headers": headers, "types": types, "rows": output}


def parse(content, suffix):
    warnings, sheets = [], []
    if suffix not in FORMATS:
        raise ValueError("Adjunta XLSX, CSV, PDF, DOCX o PPTX.")
    archive = office_archive(content, suffix) if suffix in {".xlsx", ".docx", ".pptx"} else None
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > MAX_SHEETS:
                    raise ValueError("El libro supera 20 hojas.")
                formula_count = 0
                for page in book.worksheets:
                    if (page.max_row or 0) > MAX_ROWS + 1 or (page.max_column or 0) > MAX_COLUMNS:
                        raise ValueError("Una hoja supera 50,000 filas o 80 columnas. Reduce el rango guardado, incluido su formato vacío.")
                    def values():
                        nonlocal formula_count
                        for row in page.iter_rows():
                            result = []
                            for cell in row:
                                formula_count += cell.data_type == "f"
                                value = cell.value
                                if isinstance(value, (int, float)) and re.fullmatch("0{2,20}", cell.number_format or ""):
                                    value = format(int(value), f"0{len(cell.number_format)}d")
                                result.append(value)
                            yield result
                    sheets.append(sheet(page.title, values()))
                if formula_count:
                    warnings.append("formulas_preserved_as_text")
            finally:
                book.close()
        elif suffix == ".csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")
                warnings.append("encoding_windows_1252")
            csv.field_size_limit(MAX_CELL)
            try:
                dialect = csv.Sniffer().sniff(text[:32000], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            # CSV has no native cell types. Preserve textual identifiers and infer
            # plain finite decimals only when the entire column is unambiguous.
            page = sheet("CSV", csv.reader(StringIO(text, newline=""), dialect=dialect, strict=True))
            for index, header in enumerate(page["headers"]):
                values = [row[index] for row in page["rows"] if row[index] not in (None, "")]
                if normalize(header) in IDENTIFIERS or normalize(header).endswith("id"):
                    continue
                if values and all(re.fullmatch(r"[+-]?(?:0|[1-9]\d*)(?:\.\d+)?", value) and len(value.replace(".", "").lstrip("+-")) <= 15 for value in values):
                    page["types"][index] = "number"
                    for row in page["rows"]:
                        if row[index] not in (None, ""):
                            row[index] = int(row[index]) if "." not in row[index] else float(row[index])
            sheets.append(page)
        elif suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("El archivo no contiene un PDF válido.")
            from pypdf import PdfReader, filters
            filters.ZLIB_MAX_OUTPUT_LENGTH = 16 * 1024 * 1024
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("Adjunta un PDF sin contraseña.")
            if len(reader.pages) > 200:
                raise ValueError("El PDF supera 200 páginas.")
            pages = [["Page", "Text"]]
            for number, page in enumerate(reader.pages, 1):
                pages.append([number, page.extract_text() or ""])
            sheets.append(sheet("Pages", pages))
            if not any(row[1].strip() for row in pages[1:]):
                warnings.append("pdf_requires_ocr")
            elif any(not row[1].strip() for row in pages[1:]):
                warnings.append("pdf_some_pages_without_text")
        elif suffix == ".docx":
            body = ET.fromstring(archive.read("word/document.xml")).find(f"{W}body")
            paragraphs = [["Paragraph", "Text"]]
            for child in body:
                if child.tag == W + "p":
                    text = "".join(node.text or "" for node in child.iter(W + "t"))
                    if text.strip():
                        paragraphs.append([len(paragraphs), text])
                elif child.tag == W + "tbl":
                    rows = [["".join(node.text or "" for node in cell.iter(W + "t")) for cell in row.findall(W + "tc")]
                            for row in child.findall(W + "tr")]
                    sheets.append(sheet(f"Table {len(sheets) + 1}", rows))
            sheets.insert(0, sheet("Text", paragraphs))
            warnings.append("text_and_tables_only")
        else:
            slides = [["Slide", "Text"]]
            # Presentation order comes from relationships, not lexical filenames.
            presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
            relationships = {node.attrib["Id"]: node.attrib["Target"] for node in ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))}
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            for slide in presentation.findall("p:sldIdLst/p:sldId", ns):
                rid = slide.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                target = relationships[rid]
                path = target.lstrip("/") if target.startswith("/") else "ppt/" + target
                if path not in archive.namelist():
                    raise ValueError("Una diapositiva tiene una relación no admitida.")
                root = ET.fromstring(archive.read(path))
                text = "\n".join("".join(node.text or "" for node in paragraph.iter(A + "t")) for paragraph in root.iter(A + "p"))
                slides.append([len(slides), text])
            sheets.append(sheet("Slides", slides))
            warnings.append("text_only_images_not_analyzed")
        if len(sheets) > MAX_SHEETS:
            raise ValueError("El archivo supera 20 hojas o tablas.")
        cells = sum(len(page["rows"]) * len(page["headers"]) for page in sheets)
        text_size = sum(len(value) for page in sheets for row in page["rows"] for value in row if isinstance(value, str))
        if cells > MAX_CELLS or text_size > MAX_TEXT:
            raise ValueError("El archivo supera el límite de celdas o texto; no se ha analizado parcialmente.")
        return {"sheets": sheets, "warnings": warnings, "complete": True}
    finally:
        if archive:
            archive.close()


def main():
    # Linux production: hard address-space/CPU bounds. Windows uses the same
    # byte/structure limits plus the parent's process deadline.
    if sys.platform != "win32":
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1024 ** 3, 1024 ** 3))
        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    try:
        content = Path(sys.argv[1]).read_bytes()
        result = parse(content, sys.argv[2])
        print(json.dumps({"ok": True, "data": result}, ensure_ascii=False, allow_nan=False))
    except Exception as error:
        known = isinstance(error, (ValueError, csv.Error, zipfile.BadZipFile))
        print(json.dumps({"ok": False, "error": str(error) if known else "No se pudo leer la estructura del archivo."}))


if __name__ == "__main__":
    main()
