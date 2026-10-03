"""Create a private, read-only analytics snapshot; never publish Excel as ASC runs."""
import argparse
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from app.modules.engine import parse_period
from app.modules.reporting.analytics import FIELDS
from app.persistence.business import column_name

NUMERIC_PREFIXES = ('valor_', 'total_', 'cantidad_', 'precio_', 'peso_', 'tasa_', 'importe_', 'titulos_')
NUMERIC_CODES = {'tipo_cambio', 'subdivision_fraccion', 'secuencia_fraccion', 'secuencia_observacion',
                 'tipo_pedimento', 'tipo_operacion', 'semaforo_fiscal', 'numero_seleccion', 'consecutivo_remesa',
                 'tipo_tasa', 'forma_pago', 'clave_contribucion', 'unidad_medida', 'unidad_medida_comercial',
                 'unidad_medida_tarifa', 'clave_vinculacion', 'metodo_valorizacion', 'destino_mercancia'}


def canonical_value(key, value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat()
    text = str(value).strip()
    if key in NUMERIC_CODES or key.startswith(NUMERIC_PREFIXES):
        try:
            parsed = Decimal(text)
            if parsed.is_finite():
                return str(parsed.normalize())
        except InvalidOperation:
            pass
    return text


def read_workbook(path):
    tables = {code: [] for code in FIELDS}
    periods = {}
    hashes = {}
    with path.open('rb') as stream:
        sha = hashlib.file_digest(stream, 'sha256').hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            code = sheet.title.split()[0].lower()
            if code not in {'sel', 'inci', 'resumen'} and not code.isdigit():
                continue
            iterator = sheet.iter_rows(values_only=True)
            headers = [column_name(str(h)) for h in next(iterator, ())]
            if 'periodo' not in headers:
                continue
            sheet_periods = set()
            for row in iterator:
                if not any(v is not None for v in row):
                    continue
                record = dict(zip(headers, row))
                year, month, name = parse_period(str(record['periodo']))
                p = periods.setdefault(name, {'year': year, 'month': month, 'name': name, 'rows': 0, 'tables': {},
                                               'warnings': None, 'errors': None, 'files': None})
                p['rows'] += 1
                p['tables'][code] = p['tables'].get(code, 0) + 1
                sheet_periods.add(name)
                canonical = {k: canonical_value(k, v)
                             for k, v in record.items() if k not in {'archivo_origen', 'folio_origen'}}
                digest = hashlib.sha256(json.dumps(canonical, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                hashes.setdefault((name, code), Counter())[digest] += 1
                if code in FIELDS:
                    tables[code].append({f: record.get(f) for f in FIELDS[code]} | {'year': year, 'month': month})
            # Empty registered sheets mean zero rows, not an absent dataset.
            if not sheet_periods:
                for p in periods.values():
                    p['tables'].setdefault(code, 0)
        for p in periods.values():
            for sheet in workbook:
                code = sheet.title.split()[0].lower()
                if code in FIELDS:
                    p['tables'].setdefault(code, 0)
    finally:
        workbook.close()
    return {'periods': list(periods.values()), 'tables': tables, 'sources': [{'file': path.name, 'sha256': sha}]}, hashes


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--compare', type=Path)
    parser.add_argument('--scope', required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    data, hashes = read_workbook(args.workbook)
    data.update(schemaVersion=1, scopeId=args.scope, reconciliation=None)
    if args.compare:
        other, other_hashes = read_workbook(args.compare)
        differences = [{'period': p, 'table': c} for (p, c), rows in other_hashes.items() if hashes.get((p, c)) != rows]
        data['reconciliation'] = {'file': args.compare.name, 'matched': not differences,
                                  'comparedTables': len(other_hashes), 'differences': differences}
        data['sources'] += other['sources']
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix('.tmp')
    temporary.write_text(json.dumps(data, ensure_ascii=False, default=str, separators=(',', ':')), encoding='utf-8')
    temporary.replace(args.output)
    print(json.dumps({'periods': len(data['periods']), 'reconciliation': data['reconciliation']}))


if __name__ == '__main__':
    main()
