"""Read-only functional XLSX comparison for DataStage migration regression.

Usage from backend with the test dependencies installed:
    python -m app.cli.compare_workbooks generated.xlsx reference.xlsx --report comparison.json

Exit codes: 0 equivalent under the reported policy; 1 differences; 2 input or I/O error.
No formulas are calculated. Empty strings and absent cells are equivalent blanks.
Cell values are omitted from reports unless --include-values is explicitly requested.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, time, timedelta
import hashlib
from itertools import zip_longest
import json
from pathlib import Path
import sys
from uuid import uuid4
from zipfile import BadZipFile

from app.modules.engine.catalog import normalize_header


def _digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _cell(cell) -> tuple[str, object]:
    value = cell.value if cell is not None else None
    if value is None or value == "":
        return "blank", None
    if getattr(cell, "data_type", None) == "f":
        return "formula", str(value)
    if getattr(cell, "data_type", None) == "e":
        return "error", str(value)
    if isinstance(value, datetime):
        return "datetime", value.isoformat(timespec="microseconds")
    if isinstance(value, date):
        return "datetime", datetime.combine(value, time()).isoformat(timespec="microseconds")
    if isinstance(value, time):
        return "time", value.isoformat(timespec="microseconds")
    if isinstance(value, timedelta):
        return "duration", value.total_seconds()
    if isinstance(value, bool):
        return "boolean", value
    if isinstance(value, (int, float)):
        return "number", value
    return "text", str(value)


def _row_extent(row: tuple) -> int:
    return max((index for index, cell in enumerate(row, 1) if _cell(cell)[0] != "blank"), default=0)


def _sheet_rows(worksheet):
    # Recompute dimensions from the XML stream instead of trusting a stale range
    # marker; only populated cells affect the reported row/column counts.
    worksheet.reset_dimensions()
    return worksheet.iter_rows()


def compare_workbooks(generated: Path | str, reference: Path | str, *, max_differences: int = 100,
                      include_values: bool = False, max_value_characters: int = 120) -> dict:
    if not 1 <= max_differences <= 100_000:
        raise ValueError("max_differences debe estar entre 1 y 100000.")
    if not 1 <= max_value_characters <= 10_000:
        raise ValueError("max_value_characters debe estar entre 1 y 10000.")
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Instala las dependencias de regresión: python -m pip install -e './backend[test]'") from exc

    generated, reference = Path(generated).resolve(), Path(reference).resolve()
    for path in (generated, reference):
        if not path.is_file():
            raise ValueError(f"No se encontró el archivo: {path}")
        if path.suffix.lower() != ".xlsx":
            raise ValueError("La comparación admite exclusivamente archivos .xlsx.")
    samples, categories, sheet_summaries = [], Counter(), []
    difference_count = 0

    def bounded(value):
        if isinstance(value, str) and len(value) > max_value_characters:
            return {"preview": value[:max_value_characters], "length": len(value), "truncated": True}
        if isinstance(value, list):
            return {"count": len(value), "preview": [bounded(v) for v in value[:20]], "truncated": len(value) > 20}
        return value

    def difference(kind, *, sheet=None, coordinate=None, expected=None, actual=None,
                   expected_type=None, actual_type=None):
        nonlocal difference_count
        difference_count += 1
        categories[kind] += 1
        if len(samples) >= max_differences:
            return
        record = {"kind": kind}
        if sheet is not None:
            record["sheet"] = sheet
        if coordinate is not None:
            record["cell"] = coordinate
        if expected_type is not None:
            record.update(expectedType=expected_type, actualType=actual_type)
        if include_values:
            record.update(expected=bounded(expected), actual=bounded(actual))
        samples.append(record)

    left = right = None
    try:
        left = load_workbook(generated, read_only=True, data_only=False, keep_links=False)
        right = load_workbook(reference, read_only=True, data_only=False, keep_links=False)
        if left.sheetnames != right.sheetnames:
            difference("sheet_order", expected=right.sheetnames, actual=left.sheetnames)
        for name in right.sheetnames:
            if name not in left.sheetnames:
                difference("missing_sheet", sheet=name, expected=name)
        for name in left.sheetnames:
            if name not in right.sheetnames:
                difference("extra_sheet", sheet=name, actual=name)
        for name in right.sheetnames:
            if name not in left.sheetnames:
                continue
            left_rows, right_rows = _sheet_rows(left[name]), _sheet_rows(right[name])
            stats = {side: {"dataRows": 0, "populatedDataRows": 0, "columns": 0,
                            "headerColumns": 0, "formulaCells": 0} for side in ("generated", "reference")}
            reference_headers = []
            for row_number, (left_row, right_row) in enumerate(zip_longest(left_rows, right_rows, fillvalue=()), 1):
                if row_number == 1:
                    reference_headers = [normalize_header(cell.value) for cell in right_row]
                for side, row in (("generated", left_row), ("reference", right_row)):
                    extent = _row_extent(row)
                    stats[side]["columns"] = max(stats[side]["columns"], extent)
                    stats[side]["formulaCells"] += sum(_cell(cell)[0] == "formula" for cell in row)
                    if row_number == 1:
                        stats[side]["headerColumns"] = extent
                    elif extent:
                        stats[side]["dataRows"] = row_number - 1
                        stats[side]["populatedDataRows"] += 1
                for column, (left_cell, right_cell) in enumerate(zip_longest(left_row, right_row), 1):
                    actual_type, actual = _cell(left_cell)
                    expected_type, expected = _cell(right_cell)
                    if (actual_type, actual) == (expected_type, expected):
                        continue
                    if row_number == 1:
                        kind = "header"
                    elif normalize_header(name) == "controlproceso":
                        kind = "control_cell"
                    elif column <= len(reference_headers) and reference_headers[column - 1] in {"pedimentocompleto", "pedimentounificado"}:
                        kind = "pedimento_cell"
                    elif actual_type != expected_type:
                        kind = "cell_type"
                    elif actual_type in {"datetime", "time", "duration"}:
                        kind = "date_value"
                    else:
                        kind = "cell_value"
                    difference(kind, sheet=name, coordinate=f"{get_column_letter(column)}{row_number}",
                               expected=expected, actual=actual, expected_type=expected_type, actual_type=actual_type)
            for metric in ("dataRows", "columns", "headerColumns"):
                if stats["generated"][metric] != stats["reference"][metric]:
                    difference("sheet_" + metric, sheet=name,
                               expected=stats["reference"][metric], actual=stats["generated"][metric])
            sheet_summaries.append({"sheet": name, **stats})
        sheet_names = {"generated": left.sheetnames, "reference": right.sheetnames}
    finally:
        if left is not None:
            left.close()
        if right is not None:
            right.close()

    return {
        "schemaVersion": 1, "match": difference_count == 0,
        "generated": {"path": str(generated), "sha256": _digest(generated), "sheets": sheet_names["generated"]},
        "reference": {"path": str(reference), "sha256": _digest(reference), "sheets": sheet_names["reference"]},
        "comparisonPolicy": {
            "values": "Comparación exacta por posición; texto conserva ceros iniciales; números enteros/decimales equivalentes por valor.",
            "blanks": "Celdas ausentes y cadenas vacías equivalentes; filas/columnas finales vacías ignoradas.",
            "dates": "Fechas Excel nativas comparadas por fecha/hora interpretada; una fecha en texto conserva tipo texto.",
            "formulas": "Se compara la expresión; no se calcula ni se compara su caché de resultados.",
            "formatting": "No se comparan estilos, dimensiones visuales ni objetos; el formato determina la interpretación nativa de fecha.",
            "sampling": "Se revisan todas las celdas; solo se limita el detalle retenido.",
            "valuesIncluded": include_values,
        },
        "differenceCount": difference_count, "differencesByKind": dict(sorted(categories.items())),
        "sampleLimit": max_differences, "samplesTruncated": difference_count > len(samples),
        "differences": samples, "sheetSummaries": sheet_summaries,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Compara dos Excel DataStage sin modificarlos.")
    parser.add_argument("generated", type=Path, help="Excel generado por el motor nuevo")
    parser.add_argument("reference", type=Path, help="Excel de referencia generado con las mismas fuentes ASC")
    parser.add_argument("--report", type=Path, help="Ruta local .json del informe")
    parser.add_argument("--max-differences", type=int, default=100, help="Máximo de diferencias detalladas; se cuentan todas")
    parser.add_argument("--include-values", action="store_true", help="Incluir valores acotados únicamente en el informe local")
    parser.add_argument("--max-value-characters", type=int, default=120)
    args = parser.parse_args(argv)
    if args.include_values and args.report is None:
        parser.error("--include-values requiere --report; los valores no se muestran en stdout.")
    try:
        report_path = args.report.resolve() if args.report else None
        if report_path:
            if report_path.suffix.lower() != ".json":
                raise ValueError("El informe debe guardarse con extensión .json.")
            for original in (args.generated.resolve(), args.reference.resolve()):
                if report_path == original or (report_path.exists() and original.exists() and report_path.samefile(original)):
                    raise ValueError("El informe no puede reemplazar un archivo comparado.")
        report = compare_workbooks(args.generated, args.reference, max_differences=args.max_differences,
                                   include_values=args.include_values, max_value_characters=args.max_value_characters)
        if report_path:
            report_path.parent.mkdir(parents=True, exist_ok=True)
            temporary = report_path.with_name(f".{report_path.name}.{uuid4().hex}.tmp")
            try:
                temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                temporary.replace(report_path)
            finally:
                temporary.unlink(missing_ok=True)
        print(json.dumps({"match": report["match"], "differenceCount": report["differenceCount"],
                          "differencesByKind": report["differencesByKind"],
                          "sampledDifferences": len(report["differences"]),
                          "report": str(report_path) if report_path else None}, ensure_ascii=False))
        return 0 if report["match"] else 1
    except (OSError, ValueError, RuntimeError, BadZipFile, KeyError, SyntaxError, TypeError) as exc:
        print(json.dumps({"error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
