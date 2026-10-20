"""Supplier-only reset, authoritative templates and inclusive September cutoff."""
import io
import json
from uuid import uuid4

import openpyxl
import pytest
from openpyxl.styles import Font
from sqlalchemy import func, select

from app.modules.identity.auth import Principal, current_principal
from app.modules.suppliers import service
from app.persistence.models import Period, Run, StoredDocument
from app.persistence.suppliers import Supplier
from test_api_worker import system as system, upload, drain
from test_supplier_monthly_automatic import entries

NATIONAL = ["RFC", "Operaciones virtuales", "Estatus"]
FOREIGN = ["Nombre o Razón social", "Apellido Paterno", "Apellido Materno", "País ",
           "Domicilio", "Cliente/Proveedor", "Estatus"]


def original(headers, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers)
    sheet["A1"].font = Font(name="Arial", bold=True, color="FF003344")
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def import_originals(client):
    reply = client.post("/api/v1/suppliers/import", files={
        "national": ("national.xlsx", original(NATIONAL, [["ABC010101AB1", "No", "Sin modificación"]])),
        "foreign": ("foreign.xlsx", original(FOREIGN, [
            ["BASE VENDOR", "", "", "USA", "BASE STREET", "Proveedor", "Sin modificación"]]))})
    assert reply.status_code == 200, reply.text


def publish(system, month, names):
    app, client, settings = system
    period = {8: "Agosto_2026", 9: "Septiembre_2026", 10: "Octubre_2026"}[month]
    response = upload(client, entries(month, names), period=period)
    assert response.status_code == 202, response.text
    drain(app, settings)
    run_id = response.json()["id"]
    assert client.get(f"/api/v1/runs/{run_id}").json()["publicationStatus"] == "published"
    return run_id


def rebuild(client, *, key=None, version=None):
    if version is None:
        version = client.get("/api/v1/suppliers/summary").json()["catalogVersion"]
    return client.post("/api/v1/suppliers/rebuild", headers={"Idempotency-Key": key or str(uuid4())},
        json={"startYear": 2026, "startMonth": 9, "expectedVersion": version,
              "confirmation": "REINICIAR LISTAS"})


def names(client):
    return {item["name"] for item in client.get("/api/v1/suppliers?type=foreign").json()["items"]}


def test_reset_preserves_originals_business_data_and_downloadable_backup(system):
    app, client, _ = system
    august = publish(system, 8, ["AUGUST ONLY"])
    september = publish(system, 9, ["SEPTEMBER NEW"])
    import_originals(client)
    assert {"AUGUST ONLY", "SEPTEMBER NEW", "BASE VENDOR"} <= names(client)
    old_report = client.get("/api/v1/suppliers/reports").json()["items"][0]["id"]
    with app.state.session_factory() as session:
        run_ids = set(session.scalars(select(Run.id)))
        period_ids = set(session.scalars(select(Period.id)))
        document_keys = set(session.scalars(select(StoredDocument.storage_key)))
    response = rebuild(client)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["startPeriod"] == {"year": 2026, "month": 9}
    assert result["national"] == 1
    assert names(client) == {"BASE VENDOR", "SEPTEMBER NEW"}
    assert client.get("/api/v1/suppliers/summary").json()["backupDocumentId"] == result["backupDocumentId"]
    reports_before = client.get("/api/v1/suppliers/reports").json()["total"]
    report_id = client.get("/api/v1/suppliers/reports").json()["items"][0]["id"]
    assert client.post("/api/v1/suppliers/sync").status_code == 200
    assert client.get("/api/v1/suppliers/reports").json()["total"] == reports_before
    current = client.get(f"/api/v1/suppliers/reports/{report_id}/foreign.xlsx")
    current_book = openpyxl.load_workbook(io.BytesIO(current.content))
    current_rows = {row[0].value: row for row in list(current_book.active.rows)[1:]}
    assert current_rows["SEPTEMBER NEW"][-1].value == "Alta"
    assert all(cell.fill.fgColor.rgb[-6:] == "FFFF00" for cell in current_rows["SEPTEMBER NEW"])
    current_book.close()
    with app.state.session_factory() as session:
        assert set(session.scalars(select(Run.id))) == run_ids
        assert set(session.scalars(select(Period.id))) == period_ids
        assert document_keys <= set(session.scalars(select(StoredDocument.storage_key)))
    assert all(app.state.storage.resolve(key).is_file() for key in document_keys)
    backup = client.get(f"/api/v1/suppliers/backups/{result['backupDocumentId']}")
    assert backup.status_code == 200
    contents = json.loads(backup.content)
    assert any(row["name"] == "AUGUST ONLY" for row in contents["tables"]["suppliers"])
    assert "processing_runs" not in contents["tables"]
    assert client.get(f"/api/v1/suppliers/reports/{old_report}/foreign.xlsx").status_code == 404
    assert client.get(f"/api/v1/runs/{august}").status_code == 200
    assert client.get(f"/api/v1/runs/{september}").status_code == 200
    pair = client.get("/api/v1/suppliers/reports").json()["items"][0]
    report = client.get(f"/api/v1/suppliers/reports/{pair['id']}/foreign.xlsx")
    book = openpyxl.load_workbook(io.BytesIO(report.content))
    assert book.active["A1"].font.name == "Arial"
    assert book.active["A1"].font.bold
    book.close()


def test_cutoff_is_persistent_and_earlier_months_cannot_return(system):
    _, client, _ = system
    publish(system, 8, ["AUGUST ONLY"])
    publish(system, 9, ["SEPTEMBER NEW"])
    import_originals(client)
    assert rebuild(client).status_code == 200
    assert client.post("/api/v1/suppliers/sync").status_code == 200
    august = publish(system, 8, ["ANOTHER OLD VENDOR"])
    assert names(client) == {"BASE VENDOR", "SEPTEMBER NEW"}
    assert client.get(f"/api/v1/suppliers/runs/{august}/report").json()["status"] == "unavailable"
    import_originals(client)
    assert client.get("/api/v1/suppliers/summary").json()["startPeriod"] == {"year": 2026, "month": 9}
    october = publish(system, 10, ["BASE VENDOR", "OCTOBER NEW"])
    state = client.get(f"/api/v1/suppliers/runs/{october}/report").json()
    assert state["status"] == "ready", state
    content = client.get(f"/api/v1/suppliers/reports/{state['report']['id']}/foreign.xlsx").content
    book = openpyxl.load_workbook(io.BytesIO(content))
    rows = {row[0].value: row for row in list(book.active.rows)[1:]}
    assert set(rows) == {"BASE VENDOR", "SEPTEMBER NEW", "OCTOBER NEW"}
    assert all(c.fill.fgColor.rgb[-6:] == "FFFF00" for c in rows["OCTOBER NEW"])
    assert rows["BASE VENDOR"][-1].value == "Sin modificación"
    book.close()


def test_rebuild_retry_is_idempotent_and_stale_versions_do_not_delete(system):
    app, client, _ = system
    import_originals(client)
    version = client.get("/api/v1/suppliers/summary").json()["catalogVersion"]
    key = str(uuid4())
    first = rebuild(client, key=key, version=version)
    assert first.status_code == 200, first.text
    retry = rebuild(client, key=key, version=version)
    assert retry.status_code == 200
    assert retry.json()["backupDocumentId"] == first.json()["backupDocumentId"]
    assert rebuild(client, version=version).status_code == 409
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(StoredDocument)
            .where(StoredDocument.kind == "supplier_backup")) == 1
    assert names(client) == {"BASE VENDOR"}


def test_reset_requires_admin_originals_and_does_not_change_other_scopes(system):
    app, client, _ = system
    assert rebuild(client).status_code == 409
    import_originals(client)
    with app.state.session_factory.begin() as session:
        row = Supplier(scope_id="other", kind="foreign", name="OTHER SCOPE", name_key="OTHERSCOPE",
                       data_json="{}", tax_id="", country="USA")
        session.add(row)
        session.flush()
        other_id = row.id
    app.dependency_overrides[current_principal] = lambda: Principal(
        "reader", "Reader", frozenset({"Reader"}), "local")
    assert rebuild(client, version=0).status_code == 403
    app.dependency_overrides.clear()
    assert rebuild(client).status_code == 200
    with app.state.session_factory() as session:
        assert session.get(Supplier, other_id).name == "OTHER SCOPE"


def test_backup_is_admin_only_scoped_and_does_not_open_general_document_downloads(system):
    app, client, _ = system
    import_originals(client)
    result = rebuild(client).json()
    document_id = result["backupDocumentId"]
    assert client.get(f"/api/v1/documents/{document_id}/download").status_code == 403
    for roles in ({"Reader"}, {"Operator"}):
        user = Principal("user", "User", frozenset(roles), "local")
        app.dependency_overrides[current_principal] = lambda: user
        assert client.get(f"/api/v1/suppliers/backups/{document_id}").status_code == 403
    app.dependency_overrides[current_principal] = lambda: Principal(
        "other-admin", "Other admin", frozenset({"Admin"}), "other")
    assert client.get(f"/api/v1/suppliers/backups/{document_id}").status_code == 404
    app.dependency_overrides.clear()


def test_wrong_confirmation_and_invalid_period_do_not_change_catalog(system):
    _, client, _ = system
    import_originals(client)
    version = client.get("/api/v1/suppliers/summary").json()["catalogVersion"]
    for extra in ({"confirmation": "YES"}, {"startMonth": 0}, {"startYear": "2026"}):
        body = {"startYear": 2026, "startMonth": 9, "expectedVersion": version,
                "confirmation": "REINICIAR LISTAS", **extra}
        response = client.post("/api/v1/suppliers/rebuild",
            headers={"Idempotency-Key": str(uuid4())}, json=body)
        assert response.status_code == 422
        assert names(client) == {"BASE VENDOR"}
        assert client.get("/api/v1/suppliers/summary").json()["catalogVersion"] == version


def test_backup_failure_leaves_the_catalog_reports_and_version_intact(system, monkeypatch):
    app, client, _ = system
    publish(system, 8, ["AUGUST ONLY"])
    import_originals(client)
    before = client.get("/api/v1/suppliers/summary").json()
    reports = client.get("/api/v1/suppliers/reports").json()

    def fail_backup(*args, **kwargs):
        raise OSError("Storage unavailable")

    monkeypatch.setattr(app.state.storage, "put_bytes", fail_backup)
    with pytest.raises(OSError, match="Storage unavailable"):
        rebuild(client)
    assert client.get("/api/v1/suppliers/summary").json() == before
    assert client.get("/api/v1/suppliers/reports").json() == reports
    assert names(client) == {"BASE VENDOR", "AUGUST ONLY"}


def test_private_historical_import_obeys_the_same_cutoff(system):
    app, client, _ = system
    import_originals(client)
    assert rebuild(client).status_code == 200
    source = {"file": "history.xlsx", "sha256": "d" * 64}
    tables = {"501": [], "505": []}
    for number, period, supplier in (
            ("0000001", "Agosto_2026", "OLD HISTORICAL VENDOR"),
            ("0000002", "Septiembre_2026", "NEW HISTORICAL VENDOR"),
            ("0000003", "Histórico inicial", "UNDATED VENDOR")):
        operation = {"Periodo": period, "Patente": "0036", "Pedimento": number,
                     "SeccionAduanera": "240", "ArchivoOrigen": "history.asc",
                     "PedimentoCompleto": "26 24 0036 " + number}
        tables["501"].append({**operation, "Rfc": "XYZ010101XY1", "TipoOperacion": "1",
                              "ClaveDocumento": "A1", "source": {**source, "row": len(tables["501"]) + 2}})
        tables["505"].append({**operation, "PaisFacturacion": "USA", "IndentFiscalProveedor": "",
                              "ProveedorMercancia": supplier,
                              "source": {**source, "row": len(tables["505"]) + 2}})
    with app.state.session_factory.begin() as session:
        service.ingest_baseline(session, Principal("admin", "Admin", frozenset({"Admin"}), "local"),
            {"baseline": {"source": source, "tables": tables}})
    assert names(client) == {"BASE VENDOR", "NEW HISTORICAL VENDOR"}
