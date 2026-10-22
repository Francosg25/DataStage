"""Corporate assistant: real SQL data, complete calculations and real files; Azure mocked."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
import json
from sqlalchemy import func, select

from app.core.config import Settings
from app.core.errors import ApplicationError
from app.main import create_app
from app.modules.foundry import validate_tool_arguments
from app.modules.foundry.artifacts import create_excel
from app.modules.foundry.data_queries import describe_data, query_data
from app.modules.identity.auth import Principal, current_principal
from app.persistence.business import BUSINESS_TABLES, HEADER_COLUMNS, normalize_header
from app.persistence.models import AgentArtifact, AgentMessage, AgentMessageDetail, AuditEvent, Conversation, Period, ProcessingFile, ProcessingTable, Run, StoredDocument

NAMES = {5: "Mayo", 6: "Junio", 7: "Julio"}
READER = Principal("analyst", "Analyst", frozenset({"Reader"}), "authorized")


@pytest.fixture
def system(tmp_path):
    app = create_app(Settings(_env_file=None, environment="test", database_url="sqlite://",
                             document_root=tmp_path / "documents", auto_annual=False, foundry_enabled=True,
                             foundry_project_endpoint="https://example.services.ai.azure.com/api/projects/test"))
    app.dependency_overrides[current_principal] = lambda: READER
    with TestClient(app) as client:
        yield app, client


def seed(session, code, rows, *, month=5, scope="authorized", published=True):
    period = Period(scope_id=scope, year=2025, month=month, name=f"{NAMES[month]}_2025", version=2)
    session.add(period)
    session.flush()
    run = Run(scope_id=scope, period_id=period.id, year=2025, kind="monthly", input_fingerprint=uuid4().hex * 2,
              publication_status="published" if published else "pending", status="completed")
    session.add(run)
    session.flush()
    period.active_run_id = run.id
    file = ProcessingFile(run_id=run.id, period_id=period.id, file_name="test.asc", table_code=code,
                          ordinal=1, sha256="0" * 64)
    session.add(file)
    session.flush()
    session.add(ProcessingTable(run_id=run.id, table_code=code, sheet_name=code, row_count=len(rows),
                                headers_json=json.dumps(list(rows[0]) if rows else [])))
    for index, values in enumerate(rows):
        business = {"period_id": period.id, "processing_run_id": run.id, "processing_file_id": file.id,
                    "archivo_origen": "test.asc", "folio_origen": "0001", "table_code": code,
                    "source_row_number": index + 1, "original_dates_json": "{}"}
        business.update({HEADER_COLUMNS[code][normalize_header(key)]: value for key, value in values.items()})
        session.execute(BUSINESS_TABLES[code].insert().values(**business))
    session.flush()
    return run.id


def specification(code="505", columns=None, **updates):
    return {"tableCode": code, "periods": ["Mayo_2025"], "columns": columns if columns is not None else ["ProveedorMercancia", "ValorDolares"],
            "filters": [], "groupBy": [], "metrics": [], "orderBy": [], **updates}


def request_chat(client, monkeypatch, tool, arguments, *, language="es", conversation_id=None):
    seen = {}

    def model(self, history, execute_tool, language="es"):
        seen["history"] = history
        result = execute_tool(tool, arguments, "call-1")
        seen.update(result)
        return {"message": "Resultado verificado", "evidence": result["evidence"],
                "attachments": [result["data"]["attachment"]] if "attachment" in result["data"] else [],
                "usage": {"inputTokens": 100, "outputTokens": 20, "totalTokens": 120, "modelCalls": 2}}

    monkeypatch.setattr("app.modules.foundry.FoundryAgent.respond", model)
    conversation_id = conversation_id or client.post("/api/v1/agent/conversations").json()["id"]
    response = client.post(f"/api/v1/agent/conversations/{conversation_id}/messages",
                           json={"message": "Haz un reporte", "language": language})
    return conversation_id, response, seen


def test_totals_are_computed_on_all_rows_not_the_preview(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        run_id = seed(session, "505", [{"ProveedorMercancia": "Vendor", "ValorDolares": "1,000.25"}] * 125)
        result = query_data(session, READER.scope_id, specification(columns=[], groupBy=["ProveedorMercancia"],
            metrics=[{"function": "sum", "column": "ValorDolares", "alias": "Total USD"},
                     {"function": "count_rows", "column": "", "alias": "Facturas"}]))
        assert result.rows == [["Vendor", Decimal("125031.25"), 125]]
        assert result.preview(1)["matchedSourceRows"] == 125
        assert result.sources[0]["runId"] == run_id
        detail = query_data(session, READER.scope_id, specification()).preview(2)
        assert len(detail["items"]) == 2 and detail["totalResultRows"] == 125 and detail["previewLimited"]


def test_filters_grouping_distinct_and_ordering_are_real(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "557", [
            {"PedimentoCompleto": "0001", "Fraccion": "00123456", "FormaPago": "0", "ImportePago": "100.25"},
            {"PedimentoCompleto": "0001", "Fraccion": "00123456", "FormaPago": "0", "ImportePago": "200.00"},
            {"PedimentoCompleto": "0002", "Fraccion": "00999999", "FormaPago": "21", "ImportePago": "9000"}])
        result = query_data(session, READER.scope_id, specification("557", columns=[], groupBy=["Fraccion"],
            filters=[{"column": "FormaPago", "operator": "eq", "value": "0"}],
            metrics=[{"function": "sum", "column": "ImportePago", "alias": "Pago"},
                     {"function": "count_distinct", "column": "PedimentoCompleto", "alias": "Pedimentos"}],
            orderBy=[{"column": "Pago", "direction": "desc"}]))
        assert result.rows == [["00123456", Decimal("300.25"), 1]]


def test_payment_code_variants_are_the_same_category_and_missing_table_is_not_zero(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "557", [{"FormaPago": "21", "ImportePago": "10"}, {"FormaPago": "21.0", "ImportePago": "20"}])
        result = query_data(session, READER.scope_id, specification("557", columns=[], groupBy=["FormaPago"],
            filters=[{"column": "FormaPago", "operator": "eq", "value": "21"}],
            metrics=[{"function": "sum", "column": "ImportePago", "alias": "Pago"}]))
        assert result.rows == [["21", Decimal("30.0")]]
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, specification())
        assert error.value.code == "agent_table_unavailable"


def test_missing_unpublished_and_foreign_periods_are_not_silently_omitted(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "10"}])
        seed(session, "505", [{"ValorDolares": "999999"}], scope="foreign", month=6)
        seed(session, "505", [{"ValorDolares": "888888"}], published=False, month=7)
        for period in ["Junio_2025", "Julio_2025"]:
            with pytest.raises(ApplicationError, match="publicada"):
                query_data(session, READER.scope_id, specification(periods=["Mayo_2025", period]))
        assert query_data(session, READER.scope_id, specification()).rows == [[None, Decimal("10.0")]]


@pytest.mark.parametrize("changes", [
    {"tableCode": "users"}, {"columns": ["password"]},
    {"columns": ["ValorDolares", "valor_dolares"]},
    {"columns": [], "metrics": [{"function": "sum", "column": "PedimentoCompleto", "alias": "Bad"}]},
    {"orderBy": [{"column": "bad_alias", "direction": "asc"}]},
    {"periods": ["Mayo_2025", "mayo_2025"]},
])
def test_closed_catalog_and_semantic_validation(system, changes):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "10"}])
        with pytest.raises(ApplicationError):
            query_data(session, READER.scope_id, specification(**changes))


def test_currency_mixing_and_invalid_source_numbers_do_not_produce_false_totals(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "100", "ValorMonedaExtranjera": "100", "MonedaFacturacion": "USD"},
                              {"ValorDolares": "bad", "ValorMonedaExtranjera": "200", "MonedaFacturacion": "EUR"}])
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, specification(columns=[], metrics=[
                {"function": "sum", "column": "ValorMonedaExtranjera", "alias": "Wrong total"}]))
        assert error.value.code == "agent_mixed_currencies"
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, specification(columns=[], metrics=[
                {"function": "sum", "column": "ValorDolares", "alias": "USD"}]))
        assert error.value.code == "agent_invalid_source_value"


def test_source_cap_rejects_full_calculation_without_partial_total(system, monkeypatch):
    app, _ = system
    monkeypatch.setattr("app.modules.foundry.data_queries.MAX_SOURCE_ROWS", 2)
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "1"}] * 3)
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, specification())
        assert error.value.code == "agent_query_too_large"


def test_real_excel_download_types_formulas_sources_and_private_ownership(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ProveedorMercancia": '=HYPERLINK("https://bad.invalid")', "NumeroFactura": "00000123",
                              "ValorDolares": "1,234.50", "FechaPagoReal": datetime(2025, 5, 15)}])
    query = specification(columns=["NumeroFactura", "ValorDolares", "FechaPagoReal", "ProveedorMercancia"])
    arguments = {"title": "Reporte mayo", "sheets": [{"name": "Detalle", "query": query}]}
    validate_tool_arguments("create_excel", arguments)
    conversation, response, _ = request_chat(client, monkeypatch, "create_excel", arguments)
    assert response.status_code == 200, response.text
    file = response.json()["attachments"][0]
    download = client.get(file["downloadUrl"])
    assert download.status_code == 200 and download.headers["cache-control"] == "no-store"
    book = load_workbook(BytesIO(download.content))
    assert book.sheetnames == ["Detalle", "Fuentes"]
    sheet = book["Detalle"]
    assert sheet["A2"].value == "00000123" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 1234.5 and sheet["B2"].data_type == "n"
    assert sheet["C2"].value == datetime(2025, 5, 15) and sheet["C2"].is_date
    assert sheet["D2"].data_type == "s" and sheet["D2"].hyperlink is None
    assert len(sheet.tables) == 1
    assert book["Fuentes"]["C2"].value == "Mayo_2025"
    _, repeated, _ = request_chat(client, monkeypatch, "create_excel", arguments, conversation_id=conversation)
    assert repeated.json()["attachments"][0]["documentId"] == file["documentId"]
    history = client.get(f"/api/v1/agent/conversations/{conversation}/messages").json()
    assert history["items"][-1]["attachments"] == [file]
    assert history["items"][-1]["usage"]["totalTokens"] == 120
    assert client.get("/api/v1/agent/conversations").json()["items"][0]["id"] == conversation
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(AgentArtifact)) == 1
        assert session.scalar(select(func.count()).select_from(StoredDocument)) == 1
        assert session.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.action == "agent.artifact.download")) == 1
    # Same scope is insufficient: private chat files and histories belong to their author.
    app.dependency_overrides[current_principal] = lambda: Principal("other", "Other", frozenset({"Reader"}), "authorized")
    assert client.get(file["downloadUrl"]).status_code == 404
    assert client.get(f"/api/v1/agent/conversations/{conversation}/messages").status_code == 404
    assert client.get("/api/v1/agent/conversations").json()["items"] == []
    assert client.get(f"/api/v1/documents/{file['documentId']}/download").status_code == 403


def test_no_file_or_document_record_remains_when_storage_or_database_fails(system, monkeypatch):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "1"}])
        conversation = Conversation(id=str(uuid4()), scope_id=READER.scope_id, user_id=READER.id)
        session.add(conversation)
        session.flush()
        original = session.flush
        calls = 0

        def fail_flush(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("simulated database failure")
            return original(*args, **kwargs)

        monkeypatch.setattr(session, "flush", fail_flush)
        with pytest.raises(OSError):
            create_excel(session, app.state.storage, READER, conversation.id,
                         {"title": "Test", "sheets": [{"name": "Data", "query": specification()}]}, "en")
        assert list(app.state.storage.root.iterdir()) == []


def test_filters_use_values_as_data_not_executable_sql(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ProveedorMercancia": "Vendor", "ValorDolares": "10"}])
        result = query_data(session, READER.scope_id, specification(filters=[
            {"column": "ProveedorMercancia", "operator": "eq", "value": "' OR 1=1; DROP TABLE ds_505;--"}]))
        assert result.rows == []
        assert session.scalar(select(func.count()).select_from(BUSINESS_TABLES["505"])) == 1
    assert any(column["type"] == "date" for column in describe_data("505")["columns"])


def test_legacy_messages_remain_readable_and_followups_include_history(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        conversation = Conversation(id=str(uuid4()), scope_id=READER.scope_id, user_id=READER.id)
        session.add(conversation)
        session.flush()
        session.add(AgentMessage(conversation_id=conversation.id, role="user", content="Quiero mayo de 2025"))
        conversation_id = conversation.id
    _, response, captured = request_chat(client, monkeypatch, "describe_data", {"tableCode": ""}, conversation_id=conversation_id)
    assert response.status_code == 200
    assert captured["history"][0]["content"] == "Quiero mayo de 2025"
    messages = client.get(f"/api/v1/agent/conversations/{conversation_id}/messages").json()["items"]
    assert messages[0]["text"] == "Quiero mayo de 2025"
    assert messages[0]["attachments"] == []


def chat_excel(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ProveedorMercancia": "Vendor", "ValorDolares": "10"}])
    arguments = {"title": "Mayo", "sheets": [{"name": "Detalle", "query": specification()}]}
    conversation_id, response, _ = request_chat(client, monkeypatch, "create_excel", arguments)
    assert response.status_code == 200
    return conversation_id, response.json()["attachments"][0], arguments


def test_delete_chat_removes_messages_metadata_and_downloads_but_keeps_other_chats_and_source_data(system, monkeypatch):
    app, client = system
    conversation_id, file, arguments = chat_excel(system, monkeypatch)
    other, other_reply, _ = request_chat(client, monkeypatch, "create_excel", arguments)
    other_file = other_reply.json()["attachments"][0]
    with app.state.session_factory() as session:
        key = session.get(StoredDocument, file["documentId"]).storage_key
    # History management works even when the optional inference integration is disabled.
    app.state.settings.foundry_enabled = False
    assert client.delete(f"/api/v1/agent/conversations/{conversation_id}").status_code == 204
    assert not app.state.storage.resolve(key).exists()
    assert client.get(f"/api/v1/agent/conversations/{conversation_id}/messages").status_code == 404
    assert client.get(file["downloadUrl"]).status_code == 404
    assert client.post(f"/api/v1/agent/conversations/{conversation_id}/messages", json={"message": "Hello"}).status_code == 503
    assert [item["id"] for item in client.get("/api/v1/agent/conversations").json()["items"]] == [other]
    assert client.get(other_file["downloadUrl"]).status_code == 200
    with app.state.session_factory() as session:
        assert session.get(Conversation, conversation_id) is None
        assert session.get(StoredDocument, file["documentId"]) is None
        assert session.scalar(select(func.count()).select_from(AgentMessage)) == 2
        assert session.scalar(select(func.count()).select_from(AgentMessageDetail)) == 1
        assert session.scalar(select(func.count()).select_from(BUSINESS_TABLES["505"])) == 1
        event = session.scalar(select(AuditEvent).where(AuditEvent.action == "agent.conversation.delete"))
        assert event.actor == READER.id and event.scope_id == READER.scope_id
        assert json.loads(event.details_json) == {"conversationId": conversation_id, "messages": 2, "documents": 1}
        assert "Vendor" not in event.details_json
    assert client.delete(f"/api/v1/agent/conversations/{conversation_id}").status_code == 404


@pytest.mark.parametrize("scope,user", [("authorized", "other"), ("other", READER.id)])
def test_delete_chat_requires_both_author_and_scope(system, monkeypatch, scope, user):
    app, client = system
    conversation_id, file, _ = chat_excel(system, monkeypatch)
    app.dependency_overrides[current_principal] = lambda: Principal(user, "Other", frozenset({"Admin"}), scope)
    assert client.delete(f"/api/v1/agent/conversations/{conversation_id}").status_code == 404
    app.dependency_overrides[current_principal] = lambda: READER
    assert client.get(file["downloadUrl"]).status_code == 200
    assert client.get(f"/api/v1/agent/conversations/{conversation_id}/messages").json()["total"] == 2


def test_delete_failure_rolls_back_history_and_never_removes_files_before_commit(system, monkeypatch):
    from sqlalchemy.orm import Session

    app, client = system
    conversation_id, file, _ = chat_excel(system, monkeypatch)
    original = Session.commit

    def fail_commit(session):
        if any(isinstance(item, AuditEvent) and item.action == "agent.conversation.delete" for item in session.new):
            raise RuntimeError("Database unavailable")
        return original(session)

    with monkeypatch.context() as patch:
        patch.setattr(Session, "commit", fail_commit)
        with pytest.raises(RuntimeError, match="Database unavailable"):
            client.delete(f"/api/v1/agent/conversations/{conversation_id}")
    assert client.get(file["downloadUrl"]).status_code == 200
    assert client.get(f"/api/v1/agent/conversations/{conversation_id}/messages").json()["total"] == 2
    with app.state.session_factory() as session:
        assert session.get(StoredDocument, file["documentId"]).status == "available"
        assert session.scalar(select(func.count()).select_from(AgentMessageDetail)) == 1
        assert session.scalar(select(AuditEvent).where(AuditEvent.action == "agent.conversation.delete")) is None


def test_locked_chat_file_is_revoked_immediately_and_cleanup_retries_safely(system, monkeypatch):
    from pathlib import Path
    from app.modules.foundry.conversations import cleanup_deleted_artifacts

    app, client = system
    conversation_id, file, _ = chat_excel(system, monkeypatch)
    with app.state.session_factory() as session:
        path = app.state.storage.resolve(session.get(StoredDocument, file["documentId"]).storage_key)
    original = Path.unlink

    def locked(candidate, *args, **kwargs):
        if candidate == path:
            raise PermissionError("File temporarily locked")
        return original(candidate, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(Path, "unlink", locked)
        assert client.delete(f"/api/v1/agent/conversations/{conversation_id}").status_code == 204
    assert path.exists()
    assert client.get(file["downloadUrl"]).status_code == 404
    with app.state.session_factory.begin() as session:
        assert session.get(StoredDocument, file["documentId"]).status == "pending_delete"
        session.add(StoredDocument(id="operational", scope_id=READER.scope_id, kind="excel", original_name="Keep.xlsx",
                                  storage_key="keep.xlsx", sha256="0" * 64, size_bytes=4, status="pending_delete"))
    app.state.storage.resolve("keep.xlsx").write_bytes(b"keep")
    assert cleanup_deleted_artifacts(app.state.session_factory, app.state.storage) == 1
    assert cleanup_deleted_artifacts(app.state.session_factory, app.state.storage) == 0
    assert not path.exists()
    assert app.state.storage.resolve("keep.xlsx").read_bytes() == b"keep"


@pytest.mark.parametrize("tool_after_delete", [False, True])
def test_chat_deleted_during_inference_cannot_recreate_messages_or_generate_files(system, monkeypatch, tool_after_delete):
    app, client = system
    conversation_id = client.post("/api/v1/agent/conversations").json()["id"]

    def model(self, history, execute_tool, language="es"):
        assert client.delete(f"/api/v1/agent/conversations/{conversation_id}").status_code == 204
        if tool_after_delete:
            execute_tool("describe_data", {"tableCode": ""}, "late-call")
        return {"message": "Late answer", "evidence": []}

    monkeypatch.setattr("app.modules.foundry.FoundryAgent.respond", model)
    result = client.post(f"/api/v1/agent/conversations/{conversation_id}/messages", json={"message": "Question"})
    assert result.status_code == 404
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(AgentMessage)) == 0
        assert session.scalar(select(func.count()).select_from(AgentArtifact)) == 0


@pytest.mark.parametrize("language", ["es", "en"])
def test_dashboard_pdf_is_generated_and_downloaded_from_the_chat(system, monkeypatch, language):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed(session, "501", [{"PedimentoCompleto": "25 24 0036 0000009", "Patente": "0036", "Pedimento": "0000009",
                              "SeccionAduanera": "240", "TipoOperacion": "1", "TipoCambio": "20"}])
    arguments = {"source": "published", "year": 2025, "startMonth": 5, "endMonth": 5,
                 "operation": "", "customs": "", "document": "", "currency": "USD"}
    _, response, _ = request_chat(client, monkeypatch, "create_analysis_report", arguments, language=language)
    assert response.status_code == 200, response.text
    file = response.json()["attachments"][0]
    download = client.get(file["downloadUrl"])
    assert file["format"] == "pdf" and download.headers["content-type"] == "application/pdf"
    assert download.content.startswith(b"%PDF-")
    assert response.json()["evidence"][0]["sources"]


def test_multi_sheet_excel_exports_full_detail_and_calculated_summary(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ProveedorMercancia": "A", "ValorDolares": "1.25"}] * 120)
    arguments = {"title": "Full report", "sheets": [{"name": "Detail", "query": specification()},
        {"name": "Summary", "query": specification(columns=[], groupBy=["ProveedorMercancia"], metrics=[
            {"function": "sum", "column": "ValorDolares", "alias": "Total USD"}])}]}
    _, response, _ = request_chat(client, monkeypatch, "create_excel", arguments, language="en")
    assert response.status_code == 200, response.text
    file = response.json()["attachments"][0]
    book = load_workbook(BytesIO(client.get(file["downloadUrl"]).content))
    assert book["Detail"].max_row == 121
    assert book["Summary"]["B2"].value == 150 and book["Summary"]["B2"].data_type == "n"
    assert "Sources" in book.sheetnames


def test_date_filters_empty_values_and_empty_aggregates_are_explicit(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ValorDolares": "10", "FechaPagoReal": datetime(2025, 5, 1)},
                              {"ValorDolares": "20", "FechaPagoReal": datetime(2025, 5, 30)},
                              {"ValorDolares": None}])
        result = query_data(session, READER.scope_id, specification(filters=[
            {"column": "FechaPagoReal", "operator": "gte", "value": "2025-05-15"}]))
        assert result.matched_rows == 1 and result.rows[0][1] == Decimal("20.0")
        result = query_data(session, READER.scope_id, specification(columns=[], filters=[
            {"column": "ValorDolares", "operator": "gt", "value": "999"}], metrics=[
            {"function": "sum", "column": "ValorDolares", "alias": "Total"},
            {"function": "count_rows", "column": "", "alias": "Rows"}]))
        assert result.rows == [[None, 0]]


def test_approved_document_search_cannot_read_environment_files(system, monkeypatch, tmp_path):
    _, client = system
    (tmp_path / "analytics.md").write_text("FP0 significa pago en efectivo. FP21 corresponde al crédito IVA/IEPS.")
    (tmp_path / ".env").write_text("TOKEN=secret-12345")
    monkeypatch.setattr("app.modules.foundry.knowledge.DOC_ROOT", tmp_path)
    _, response, captured = request_chat(client, monkeypatch, "search_documentation", {"query": "FP0 FP21 TOKEN"})
    assert response.status_code == 200
    assert captured["data"]["items"][0]["source"] == "docs/analytics.md"
    assert "secret-12345" not in str(captured)


def test_artifact_migration_preserves_existing_conversations_and_documents(tmp_path):
    from alembic import command
    from alembic.config import Config
    from pathlib import Path
    from sqlalchemy import create_engine, inspect, text

    backend = Path(__file__).resolve().parents[1]
    config = Config(str(backend / "alembic.ini"))
    config.set_main_option("script_location", str(backend / "migrations"))
    engine = create_engine(f"sqlite:///{tmp_path / 'migration.db'}")
    with engine.begin() as connection:
        config.attributes["connection"] = connection
        command.upgrade(config, "20261015_0003")
        connection.execute(text("INSERT INTO agent_conversations (id,scope_id,user_id,created_at) VALUES ('c','s','u','2025-05-01')"))
        connection.execute(text("INSERT INTO agent_messages (id,conversation_id,role,content,created_at) VALUES ('m','c','user','Keep this history','2025-05-01')"))
        connection.execute(text("INSERT INTO stored_documents (id,scope_id,kind,original_name,storage_key,sha256,size_bytes,status,created_at) VALUES ('d','s','excel','Keep.xlsx','keep.xlsx','hash',100,'available','2025-05-01')"))
        command.upgrade(config, "head")
        assert connection.scalar(text("SELECT content FROM agent_messages WHERE id='m'")) == "Keep this history"
        assert connection.scalar(text("SELECT storage_key FROM stored_documents WHERE id='d'")) == "keep.xlsx"
        assert {"agent_artifacts", "agent_message_details"} <= set(inspect(connection).get_table_names())
        assert connection.scalar(text("SELECT version_num FROM alembic_version")) == "20261022_0004"
    engine.dispose()
