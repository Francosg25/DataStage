from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from app.modules.engine import InputFile, export_xlsx, process_annual, process_monthly


def test_excel_preserves_text_dates_readable_headers_and_filtered_ranges(tmp_path):
    result = process_monthly("Abril_2026", [InputFile("x_501.asc", "Clave|FechaPagoReal|Nota\n00001|20260423121314|=SUM(A1:A2)\n00002|invalid|https://example.com", "Abril_2026")])
    path = tmp_path / "DataStage_Abril_2026.xlsx"
    export_xlsx(result, path)
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Control_Proceso", "501 Datos generales"]
    worksheet = workbook.worksheets[1]
    assert worksheet.freeze_panes == "A2"
    assert worksheet["D2"].value == "00001"
    assert worksheet["D2"].number_format == "@"
    assert worksheet["E2"].value == datetime(2026, 4, 23, 12, 13, 14)
    assert worksheet["E2"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert worksheet["E3"].value == "invalid"
    assert worksheet["F2"].value == "=SUM(A1:A2)"
    assert worksheet["F2"].data_type == "s"
    assert worksheet["F3"].hyperlink is None
    assert worksheet["A1"].fill.fgColor.rgb == "FF145C56"
    assert worksheet["A1"].font.color.rgb == "FFFFFFFF"
    assert worksheet["A1"].font.bold
    assert worksheet["A1"].alignment.horizontal == "left"
    assert worksheet["A1"].alignment.wrap_text
    assert worksheet.row_dimensions[1].height == 32
    assert worksheet.auto_filter.ref == "A1:F3"
    assert all(not sheet.tables for sheet in workbook)
    with ZipFile(path) as archive:
        assert "<f>" not in archive.read("xl/worksheets/sheet2.xml").decode()
        assert not any(name.startswith("xl/tables/") for name in archive.namelist())


def test_header_only_sheet_has_autofilter(tmp_path):
    result = process_monthly("Abril_2026", [InputFile("x_501.asc", "Clave|Valor", "Abril_2026")])
    path = tmp_path / "empty.xlsx"
    export_xlsx(result, path)
    worksheet = openpyxl.load_workbook(path).worksheets[1]
    assert not worksheet.tables
    assert worksheet.auto_filter.ref == "A1:E1"
    assert worksheet.max_row == 1


def test_multiline_headers_have_readable_height_and_filter_space(tmp_path):
    result = {"kind": "monthly", "controlRows": [], "tables": [{
        "sheetName": "Encabezados", "headers": ["DescripcionMercancia", "Importe\noriginal"],
        "dateColumns": [], "rows": [["Articulo", "0001"]],
    }]}
    path = tmp_path / "headers.xlsx"
    export_xlsx(result, path)
    worksheet = openpyxl.load_workbook(path)["Encabezados"]
    assert worksheet.row_dimensions[1].height == 44
    assert worksheet.column_dimensions["A"].width >= len("DescripcionMercancia") + 4
    assert worksheet["B1"].value == "Importe\noriginal"
    assert worksheet["B1"].alignment.wrap_text
    assert worksheet["B2"].value == "0001"
    assert worksheet.auto_filter.ref == "A1:B2"
    assert not worksheet.tables


def test_annual_control_column_order_and_export(tmp_path):
    result = process_annual(2026, "Ene-Ago", [InputFile("x_501.asc", "A\n001", "Jan_2026")])
    path = tmp_path / "annual.xlsx"
    export_xlsx(result, path)
    workbook = openpyxl.load_workbook(path)
    worksheet = workbook.worksheets[0]
    assert [cell.value for cell in worksheet[1]] == ["Periodo", "Archivo", "Tabla", "Hoja", "Estatus", "Filas", "Columnas", "ColumnasFecha", "Mensaje"]
    for sheet in workbook:
        assert not sheet.tables
        assert sheet.auto_filter.ref == sheet.dimensions
        assert sheet.freeze_panes == "A2"
        assert all(cell.font.color.rgb == "FFFFFFFF" for cell in sheet[1])
        assert all(cell.fill.fgColor.rgb == "FF145C56" for cell in sheet[1])
    with ZipFile(path) as archive:
        assert not any(name.startswith("xl/tables/") for name in archive.namelist())


def test_export_limits_fail_before_replacing_existing_file(tmp_path):
    path = tmp_path / "output.xlsx"
    path.write_bytes(b"existing output")
    result = process_monthly("Abril_2026", [InputFile("x_501.asc", "A\n" + "x" * 32768, "Abril_2026")])
    with pytest.raises(ValueError, match="32767"):
        export_xlsx(result, path)
    assert path.read_bytes() == b"existing output"


def test_empty_result_still_exports_control(tmp_path):
    result = process_monthly("Abril_2026", [])
    path = tmp_path / "output.xlsx"
    export_xlsx(result, path)
    assert openpyxl.load_workbook(path).sheetnames == ["Control_Proceso"]


def test_export_native_measures_dates_and_counts_preserves_identifiers_monthly_and_annual(tmp_path):
    content = ("Fraccion|SecuenciaFraccion|PrecioUnitario|ValorDolares|CantidadUMComercial|"
               "TipoCambio|FechaPagoReal|TipoFecha|CodigoMercanciaProducto\n"
               "00123456|0001|123.456789|1195277.50|3.125|17.18470|20260423121314|01|00000987654321012345\n"
               "00123456|0002|0|-42.50|0|19.1754|20260424|02|PN-0002")
    files = [InputFile("x_551.asc", content, "Abril_2026")]
    for kind, result in (("monthly", process_monthly("Abril_2026", files)),
                         ("annual", process_annual(2026, "Ene-Ago", files))):
        path = tmp_path / f"typed-{kind}.xlsx"
        export_xlsx(result, path)
        book = openpyxl.load_workbook(path)
        sheet = book["551 Partidas"]
        fields = {cell.value: cell.column for cell in sheet[1]}
        for header, value in {"PrecioUnitario": 123.456789, "ValorDolares": 1195277.5,
                              "CantidadUMComercial": 3.125, "TipoCambio": 17.1847}.items():
            cell = sheet.cell(2, fields[header])
            assert cell.data_type == "n"
            assert cell.value == pytest.approx(value)
            assert cell.number_format != "@"
        assert sheet.cell(3, fields["PrecioUnitario"]).value == 0
        assert sheet.cell(3, fields["ValorDolares"]).value == -42.5
        for header, value in {"Fraccion": "00123456", "SecuenciaFraccion": "0001",
                              "TipoFecha": "01", "CodigoMercanciaProducto": "00000987654321012345"}.items():
            cell = sheet.cell(2, fields[header])
            assert cell.data_type == "s" and cell.value == value
        assert sheet.cell(2, fields["FechaPagoReal"]).value == datetime(2026, 4, 23, 12, 13, 14)
        assert sheet.cell(3, fields["FechaPagoReal"]).value == datetime(2026, 4, 24)
        assert sheet.cell(3, fields["FechaPagoReal"]).number_format == "yyyy-mm-dd"
        control = book["Control_Proceso"]
        counts = {cell.value: cell.column for cell in control[1]}
        for header in ("Filas", "Columnas", "ColumnasFecha"):
            assert control.cell(2, counts[header]).data_type == "n"


def test_numeric_export_keeps_invalid_and_overprecision_values_without_coercing_to_zero(tmp_path):
    result = process_monthly("Abril_2026", [InputFile("x_557.asc",
        "ImportePago|FormaPago|FechaPagoReal\nNaN|00|invalid\n1e309|21|invalid\n"
        "1234567890123456|00|invalid\n=SUM(A1:A2)|00|invalid\n1,234.50|00|invalid\n"
        "1.23456789012345|00|20260423", "Abril_2026")])
    path = tmp_path / "numeric-originals.xlsx"
    export_xlsx(result, path)
    sheet = openpyxl.load_workbook(path)["557 Contribuciones partida"]
    for row, original in enumerate(("NaN", "1e309", "1234567890123456", "=SUM(A1:A2)", "1,234.50"), 2):
        assert sheet[f"D{row}"].value == original and sheet[f"D{row}"].data_type == "s"
    assert sheet["D7"].data_type == "n" and sheet["D7"].value == pytest.approx(1.23456789012345)
    assert sheet["E2"].value == "00" and sheet["E2"].data_type == "s"
    assert sheet["F2"].value == "invalid" and sheet["F2"].data_type == "s"


def test_out_of_range_iso_date_remains_original_text_in_excel(tmp_path):
    result = process_monthly("Abril_2026", [InputFile("x_501.asc", "FechaPagoReal\n2200-01-01T00:00:00", "Abril_2026")])
    path = tmp_path / "invalid-date.xlsx"
    export_xlsx(result, path)
    cell = openpyxl.load_workbook(path).worksheets[1]["D2"]
    assert cell.value == "2200-01-01T00:00:00"
    assert cell.data_type == "s"


def test_empty_cells_are_not_serialized_or_merged_and_values_keep_their_columns(tmp_path):
    result = {"kind": "monthly", "controlRows": [], "tables": [{
        "sheetName": "Vacios", "headers": ["Clave", "Vacio", "Valor", "FechaPagoReal"],
        "dateColumns": ["fechapagoreal"],
        "rows": [["00001", "", "0", "2026-04-23T12:13:14"],
                 [None, "", "  ", None], ["00002", None, "42", None]],
    }]}
    path = tmp_path / "sparse.xlsx"
    export_xlsx(result, path)
    sheet = openpyxl.load_workbook(path)["Vacios"]
    assert list(sheet.values) == [
        ("Clave", "Vacio", "Valor", "FechaPagoReal"),
        ("00001", None, "0", datetime(2026, 4, 23, 12, 13, 14)),
        ("00002", None, "42", None),
    ]
    assert sheet.auto_filter.ref == "A1:D3"
    assert not sheet.merged_cells.ranges
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        coordinates = {cell.attrib["r"] for cell in xml.findall(".//s:sheetData/s:row/s:c", ns)}
        assert coordinates == {"A1", "B1", "C1", "D1", "A2", "C2", "D2", "A3", "C3"}


def test_monthly_and_annual_excel_exclude_separator_only_records_before_trace_columns(tmp_path):
    files = [InputFile("x_501.asc", "Clave|Valor\n||\n00001|0\n | |\n00002|42\n|||", "Abril_2026")]
    for kind, result in (("monthly", process_monthly("Abril_2026", files)),
                         ("annual", process_annual(2026, "Ene-Ago", files))):
        assert len(result["tables"][0]["rows"]) == 2
        assert result["controlRows"][0]["Filas"] == 2
        path = tmp_path / f"{kind}.xlsx"
        export_xlsx(result, path)
        sheet = openpyxl.load_workbook(path).worksheets[1]
        assert sheet.max_row == 3
        assert sheet.auto_filter.ref.endswith("3")
        assert sheet["D2"].value == "00001"
        assert sheet["E2"].value == "0"

