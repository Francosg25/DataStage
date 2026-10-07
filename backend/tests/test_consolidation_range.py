import io
import json
import zipfile
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import AnnualRequest
from app.modules.engine import InputFile, process_annual
from app.modules.engine.models import ConsolidationRange
from app.modules.identity.auth import Principal, current_principal
from app.persistence.models import AnnualSource, Run
from test_api_worker import ASC, drain, get_run, system, upload


RANGE = {"startYear": 2025, "startMonth": 4, "endYear": 2026, "endMonth": 3}


def consolidate(client, body=None, key=None):
    return client.post("/api/v1/annual-runs", json=body or RANGE,
                       headers={"Idempotency-Key": key or str(uuid4())})


@pytest.mark.parametrize("changes", [
    {"startYear": None}, {"endMonth": None}, {"startYear": 1899}, {"endYear": 2101},
    {"startMonth": 0}, {"endMonth": 13}, {"startMonth": True}, {"endMonth": 3.5},
    {"startYear": "2025"}, {"startYear": 2027}, {"endYear": 2025, "endMonth": 3},
    {"anio": 2025, "rangoNombre": "Ene-Mar"},
])
def test_range_request_rejects_incomplete_invalid_or_mixed_values(changes):
    with pytest.raises(ValueError):
        AnnualRequest.model_validate({**RANGE, **changes})


@pytest.mark.parametrize("bounds,expected", [
    ((2025, 4, 2026, 3), ["Abril_2025", "Diciembre_2025", "Enero_2026", "Marzo_2026"]),
    ((2026, 3, 2026, 4), ["Marzo_2026", "Abril_2026"]),
    ((2025, 12, 2025, 12), ["Diciembre_2025"]),
])
def test_engine_includes_exact_bounds_and_preserves_each_source_year(bounds, expected):
    periods = ["Abril_2026", "Marzo_2026", "Enero_2026", "Diciembre_2025", "Abril_2025", "Marzo_2025", "invalid"]
    files = [InputFile(f"{i}_501.asc", "Patente|Pedimento|SeccionAduanera|\n0036|0000009|240|", p)
             for i, p in enumerate(periods)]
    interval = ConsolidationRange(*bounds)
    result = process_annual(interval.start_year, "unused", files, period_range=interval)
    assert result["periodRange"] == interval.payload()
    assert result["rangeName"] == interval.label
    assert result["processedFiles"] == len(expected)
    assert len(result["excludedFiles"]) == len(periods) - len(expected)
    table = result["tables"][0]
    rows = [dict(zip(table["headers"], row)) for row in table["rows"]]
    assert [row["Periodo"] for row in rows] == expected
    assert all(row["PedimentoCompleto"].startswith(row["Periodo"][-2:] + " ") for row in rows)
    assert not any(w["code"] == "UNRECOGNIZED_RANGE" for w in result["warnings"])


@pytest.mark.parametrize("bounds", [(2026, 4, 2025, 4), (2026, 5, 2026, 4),
                                        (2026, 0, 2026, 4), (True, 4, 2026, 4)])
def test_domain_range_validates_independent_of_http(bounds):
    with pytest.raises(ValueError):
        ConsolidationRange(*bounds)


def test_cross_year_snapshot_export_and_reprocess(system):
    app, client, settings = system
    months = ["Marzo_2025", "Abril_2025", "Diciembre_2025", "Marzo_2026", "Abril_2026"]
    runs = {}
    for period in months:
        response = upload(client, period=period)
        drain(app, settings)
        runs[period] = get_run(client, response)["id"]

    key = str(uuid4())
    response = consolidate(client, key=key)
    assert response.status_code == 202, response.text
    drain(app, settings)
    annual = get_run(client, response)
    assert annual["status"] == "completed", annual
    assert annual["periodRange"] == RANGE
    assert annual["rangeName"] == "2025-04 - 2026-03"
    assert annual["counts"]["rows"] == 3
    included = ["Abril_2025", "Diciembre_2025", "Marzo_2026"]
    with app.state.session_factory() as session:
        run = session.get(Run, annual["id"])
        assert json.loads(run.source_manifest_json) == [{"monthlyRunId": runs[p]} for p in included]
        assert set(session.scalars(select(AnnualSource.monthly_run_id).where(
            AnnualSource.annual_run_id == run.id))) == {runs[p] for p in included}

    download = client.get(f"/api/v1/documents/{annual['exportDocumentId']}/download")
    assert download.status_code == 200
    assert "DataStage_2025-04_2026-03.xlsx" in download.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(download.content)) as zipped:
        assert not any(name.startswith("xl/tables/") for name in zipped.namelist())
    workbook = load_workbook(io.BytesIO(download.content))
    sheet = workbook.worksheets[1]
    headers = [cell.value for cell in sheet[1]]
    assert [row[headers.index("Periodo")] for row in sheet.iter_rows(min_row=2, values_only=True)] == included
    assert sheet["A1"].fill.fgColor.rgb == "FF145C56"
    assert sheet["A1"].font.color.rgb == "FFFFFFFF"
    workbook.close()

    # A changed monthly pointer must not change either a retry or a reprocess.
    replacement = upload(client, {"changed_501.asc": ASC.replace("RFCDEMO", "RFCNEW")}, period="Abril_2025")
    drain(app, settings)
    assert get_run(client, replacement)["id"] != runs["Abril_2025"]
    assert consolidate(client, key=key).json()["id"] == annual["id"]
    assert consolidate(client, {**RANGE, "startMonth": 5}, key=key).status_code == 409
    repro = client.post(f"/api/v1/runs/{annual['id']}/reprocess",
                       json={"reason": "Verificar rango original", "expectedVersion": 0},
                       headers={"Idempotency-Key": str(uuid4())})
    drain(app, settings)
    repeated = get_run(client, repro)
    assert repeated["status"] == "completed", repeated
    assert repeated["periodRange"] == RANGE
    assert repeated["counts"]["rows"] == 3
    old_data = client.get("/api/v1/data/501", params={"runId": annual["id"]}).json()
    new_data = client.get("/api/v1/data/501", params={"runId": repeated["id"]}).json()
    assert new_data == old_data
    with app.state.session_factory() as session:
        assert session.get(Run, repeated["id"]).source_manifest_json == session.get(Run, annual["id"]).source_manifest_json


def test_range_http_validation_empty_sources_and_authorization(system):
    app, client, settings = system
    assert consolidate(client, {**RANGE, "startMonth": 13}).status_code == 422
    assert consolidate(client, {**RANGE, "endYear": 2024}).status_code == 422
    assert consolidate(client).status_code == 409
    response = upload(client, period="Abril_2026")
    drain(app, settings)
    run_id = get_run(client, response)["id"]
    # Explicit sources cannot smuggle an out-of-range month into the snapshot.
    outside = consolidate(client, {**RANGE, "sourceRunIds": [run_id]})
    assert outside.status_code == 409
    assert outside.json()["title"] == "SOURCE_OUTSIDE_RANGE"
    same_month = {"startYear": 2026, "startMonth": 4, "endYear": 2026, "endMonth": 4}
    single = consolidate(client, same_month)
    drain(app, settings)
    assert get_run(client, single)["counts"]["rows"] == 1
    app.dependency_overrides[current_principal] = lambda: Principal("reader", "Reader", frozenset({"Reader"}), settings.scope_id)
    assert consolidate(client, same_month).status_code == 403
    app.dependency_overrides[current_principal] = lambda: Principal("other", "Other", frozenset({"Operator", "Reader"}), "elsewhere")
    assert consolidate(client, same_month).status_code == 409
    assert consolidate(client, {**same_month, "sourceRunIds": [run_id]}).status_code == 404


def test_direct_files_are_filtered_by_the_explicit_range(system):
    app, client, settings = system
    sources = []
    for period in ["Marzo_2025", "Abril_2025", "Marzo_2026", "Abril_2026", "invalid"]:
        response = client.post("/api/v1/source-files", data={"periodo": period},
                               files={"content": ("direct_501.asc", ASC, "text/plain")})
        assert response.status_code == 201, response.text
        sources.append({key: value for key, value in response.json().items() if key != "sha256"})
    response = consolidate(client, {**RANGE, "files": sources})
    drain(app, settings)
    annual = get_run(client, response)
    assert annual["status"] == "completed", annual
    assert annual["counts"]["processedFiles"] == 2
    assert annual["counts"]["skippedFiles"] == 3
    page = client.get("/api/v1/data/501", params={"runId": annual["id"]}).json()
    assert page["total"] == 2
    assert {row["Periodo"] for row in page["items"]} == {"Abril_2025", "Marzo_2026"}


def test_explicit_source_contract_allows_more_than_twelve_months():
    request = AnnualRequest.model_validate({**RANGE, "endYear": 2027,
                                           "sourceRunIds": [str(uuid4()) for _ in range(24)]})
    assert len(request.sourceRunIds) == 24
