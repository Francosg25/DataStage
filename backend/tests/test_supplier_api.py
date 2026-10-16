"""HTTP workflow: permissions, snapshots, downloads and optimistic edits."""
import io
from uuid import uuid4

import openpyxl
import test_api_worker

from app.modules.identity.auth import Principal, current_principal

system = test_api_worker.system


def supplier(name="Proveedor extranjero", **fields):
    return {"kind": "foreign", "name": name, "taxId": "US-0007", "country": "USA", **fields}


def create(client, **fields):
    response = client.post("/api/v1/suppliers", json=supplier(**fields))
    assert response.status_code == 201, response.text
    return response.json()


def report(client, key=None, label="Reporte de prueba"):
    response = client.post("/api/v1/suppliers/reports", json={"label": label},
                           headers={"Idempotency-Key": key or str(uuid4())})
    assert response.status_code == 201, response.text
    return response.json()


def workbook(client, report_id, kind):
    response = client.get(f"/api/v1/suppliers/reports/{report_id}/{kind}.xlsx")
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    return response.content, openpyxl.load_workbook(io.BytesIO(response.content)).active


def test_pairs_are_immutable_and_only_new_rows_are_yellow(system):
    app, client, _ = system
    create(client)
    create(client, name="Proveedor nacional", kind="national", taxId="ABC010101AB1", country="MEX",
           virtualConfirmation=True)
    key = str(uuid4())
    first = report(client, key)
    assert report(client, key)["id"] == first["id"]
    national_bytes, national = workbook(client, first["id"], "national")
    foreign_bytes, foreign = workbook(client, first["id"], "foreign")
    for sheet in (national, foreign):
        assert sheet.max_row == 2
        assert "Alta" in [c.value for c in sheet[2]]
        assert all(cell.fill.fgColor.rgb[-6:] == "FFFF00" for cell in sheet[2])
    create(client, name="Otro proveedor", taxId="US-0008")
    second = report(client)
    _, second_foreign = workbook(client, second["id"], "foreign")
    status_column = next(c.column for c in second_foreign[1] if "estatus" in str(c.value).casefold())
    statuses = [second_foreign.cell(r, status_column).value for r in range(2, second_foreign.max_row + 1)]
    assert sorted(statuses) == ["Alta", "Sin modificación"]
    for row in list(second_foreign.rows)[1:]:
        is_new = row[status_column - 1].value == "Alta"
        assert all((cell.fill.fgColor.rgb[-6:] == "FFFF00") == is_new
                   for cell in row if cell.fill.fgColor.type == "rgb")
    assert workbook(client, first["id"], "national")[0] == national_bytes
    assert workbook(client, first["id"], "foreign")[0] == foreign_bytes
    assert client.get("/api/v1/suppliers/reports").json()["total"] == 2


def test_roles_scope_validation_and_audit(system):
    app, client, _ = system
    row = create(client)
    snapshot = report(client)
    app.dependency_overrides[current_principal] = lambda: Principal("reader", "Reader", frozenset({"Reader"}), "local")
    assert client.get("/api/v1/suppliers").status_code == 200
    assert client.get(f"/api/v1/suppliers/{row['id']}").status_code == 200
    assert client.post("/api/v1/suppliers", json=supplier()).status_code == 403
    assert client.post("/api/v1/suppliers/sync").status_code == 403
    assert client.post("/api/v1/suppliers/reports", json={}, headers={"Idempotency-Key": str(uuid4())}).status_code == 403
    app.dependency_overrides[current_principal] = lambda: Principal("other", "Other", frozenset({"Operator", "Reader"}), "other")
    assert client.get("/api/v1/suppliers").json()["total"] == 0
    assert client.get(f"/api/v1/suppliers/{row['id']}").status_code == 404
    assert client.get(f"/api/v1/suppliers/reports/{snapshot['id']}/foreign.xlsx").status_code == 404
    app.dependency_overrides.clear()
    assert client.get("/api/v1/suppliers?offset=-1").status_code == 422
    assert client.post("/api/v1/suppliers", json=supplier(virtualConfirmation="false")).status_code == 422
    events = client.get("/api/v1/audit").json()["items"]
    assert any(item["action"].startswith("supplier.") and item["actor"] for item in events)


def test_national_confirmation_and_edit_version(system):
    _, client, _ = system
    fields = supplier(name="Proveedor nacional", kind="national", taxId="ABC010101AB1", country="MEX")
    missing = client.post("/api/v1/suppliers", json=fields)
    assert missing.status_code == 400
    assert missing.json()["title"] == "VIRTUAL_CONFIRMATION_REQUIRED"
    row = create(client, name="Proveedor nacional", kind="national", taxId="ABC010101AB1", country="MEX", virtualConfirmation=False)
    update = {**fields, "name": "Proveedor nacional editado", "expectedVersion": row["version"], "virtualConfirmation": True}
    edited = client.patch(f"/api/v1/suppliers/{row['id']}", json=update)
    assert edited.status_code == 200, edited.text
    assert edited.json()["virtualOperations"] is True
    assert client.patch(f"/api/v1/suppliers/{row['id']}", json=update).status_code == 409
    detail = client.get(f"/api/v1/suppliers/{row['id']}").json()
    assert detail["name"] == update["name"]
    assert client.get("/api/v1/suppliers?search=EDITADO").json()["total"] == 1
    summary = client.get("/api/v1/suppliers/summary").json()
    assert summary["catalog"] == summary["national"] == 1
