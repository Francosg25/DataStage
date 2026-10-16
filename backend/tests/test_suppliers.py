"""Supplier identity, live evidence and transactional publication regression tests."""
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.core.errors import ApplicationError
from app.modules.identity.auth import Principal
from app.modules.suppliers import service as suppliers
from app.persistence import Base, make_engine, session_factory
from app.persistence.business import BUSINESS_TABLES
from app.persistence.models import AuditEvent, Period, ProcessingFile, Run


@pytest.fixture
def domain():
    engine = make_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = session_factory(engine)
    with factory.begin() as session:
        yield session, Principal("operator", "Operator", frozenset({"Admin"}), "tenant")
    engine.dispose()


def publish(session, principal, month, invoices=(), identifiers=(), *, contributor="CON010101AA1"):
    period = session.scalar(select(Period).where(Period.scope_id == principal.scope_id, Period.month == month))
    if period is None:
        period = Period(scope_id=principal.scope_id, year=2026, month=month,
                        name=["Enero_2026", "Febrero_2026", "Marzo_2026"][month - 1])
        session.add(period)
        session.flush()
    run = Run(scope_id=principal.scope_id, period_id=period.id, year=2026,
              input_fingerprint=uuid4().hex * 2, status="completed", publication_status="published")
    session.add(run)
    session.flush()
    for ordinal, (code, rows) in enumerate((("501", [{"rfc": contributor, "clave_documento": "V1"}]),
                                           ("505", invoices), ("507", identifiers))):
        source = ProcessingFile(run_id=run.id, period_id=period.id, file_name=f"source_{code}.asc",
                                table_code=code, sha256=uuid4().hex * 2, ordinal=ordinal)
        session.add(source)
        session.flush()
        for row_number, row in enumerate(rows, 2):
            data = {"period_id": period.id, "processing_run_id": run.id, "processing_file_id": source.id,
                    "archivo_origen": source.file_name, "folio_origen": "source", "table_code": code,
                    "source_row_number": row_number, "pedimento_completo": "26 24 0036 6000001"}
            session.execute(BUSINESS_TABLES[code].insert().values(**{**data, **row}))
    period.active_run_id, period.version = run.id, period.version + 1
    session.flush()
    suppliers.sync_active_history(session, principal.scope_id, principal.id)
    return run


def invoice(name="Álamo Tecnología, S.A.", tax="TAX-100", country="USA", **kwargs):
    return {"proveedor_mercancia": name, "indent_fiscal_proveedor": tax, "pais_facturacion": country,
            "calle_proveedor": "100 Main St.", **kwargs}


def test_monthly_history_normalization_similar_review_and_saved_decision(domain):
    session, principal = domain
    publish(session, principal, 1, [invoice()])
    first = suppliers.list_suppliers(session, principal.scope_id)[0]
    publish(session, principal, 2, [invoice(name="alamo tecnologia SA", tax="TAX100", country="United States")])
    assert len(suppliers.list_suppliers(session, principal.scope_id)) == 1
    assert suppliers.list_reviews(session, principal.scope_id) == []
    publish(session, principal, 3, [invoice(name="ALAMO TECNOLOGIAS SA", tax="")])
    reviews = suppliers.list_reviews(session, principal.scope_id)
    assert len(reviews) == 1
    assert reviews[0]["reason"] == "similar_name"
    result = suppliers.resolve_review(session, principal, reviews[0]["id"], "same", first["id"],
                                      expected_version=reviews[0]["version"], reason="Misma razón social")
    assert result["id"] == first["id"]
    assert suppliers.list_reviews(session, principal.scope_id) == []
    # A replacement upload with the same spelling automatically reuses the decision.
    publish(session, principal, 3, [invoice(name="Alamo Tecnologias S.A.", tax="")])
    assert len(suppliers.list_suppliers(session, principal.scope_id)) == 1
    assert suppliers.list_reviews(session, principal.scope_id) == []
    alias = suppliers.list_equivalences(session, principal.scope_id)[0]
    suppliers.update_equivalence(session, principal, alias["id"], expected_version=alias["version"],
                                  active=False, reason="Retirar equivalencia")
    assert suppliers.list_reviews(session, principal.scope_id)[0]["reason"] == "equivalence_revoked"
    with pytest.raises(ApplicationError, match="cambió"):
        suppliers.update_equivalence(session, principal, alias["id"], expected_version=alias["version"], active=True)
    alias = suppliers.list_equivalences(session, principal.scope_id)[0]
    suppliers.update_equivalence(session, principal, alias["id"], first["id"], alias["version"], active=True)
    assert suppliers.list_reviews(session, principal.scope_id) == []


def test_virtual_or_history_no_reset_and_replaced_evidence_invalidated(domain):
    session, principal = domain
    national = invoice("Proveedor Nacional SA", "ABC010101AA1", "MEX")
    first = publish(session, principal, 1, [national], [
        {"clave_caso": "V1", "identificador_caso": "ABC010101AA1", "complemento_caso": "CN"}])
    supplier = suppliers.list_suppliers(session, principal.scope_id)[0]
    assert supplier["virtualOperations"] is True
    publish(session, principal, 2, [national])
    assert suppliers.get_supplier(session, principal.scope_id, supplier["id"])["virtualOperations"] is True
    # Replacing the month removes its virtual support, without deleting the supplier.
    publish(session, principal, 1, [national])
    detail = suppliers.get_supplier(session, principal.scope_id, supplier["id"])
    assert detail["virtualOperations"] is None
    assert not next(e for e in detail["evidence"] if e["runId"] == first.id)["active"]
    updated = suppliers.update_supplier(session, principal, supplier["id"],
                                         {"virtualConfirmation": True, "reason": "Confirmación de historial"}, detail["version"])
    suppliers.retire_runs(session, principal.scope_id, [e["runId"] for e in updated["evidence"]])
    assert suppliers.get_supplier(session, principal.scope_id, supplier["id"])["virtualOperations"] is True


def test_immex_and_contributor_rfc_never_become_supplier_rfc(domain):
    session, principal = domain
    publish(session, principal, 1, [invoice("Nacional sin evidencia", "ABC010101AA1", "MEX")], [
        {"clave_caso": "V1", "identificador_caso": "42832006", "complemento_caso": "IM"}])
    review = suppliers.list_reviews(session, principal.scope_id)[0]
    assert review["reason"] == "virtual_confirmation"
    assert review["candidate"]["taxId"] == "ABC010101AA1"
    with pytest.raises(ApplicationError, match="Confirma"):
        suppliers.resolve_review(session, principal, review["id"], "new", expected_version=review["version"])
    item = suppliers.resolve_review(session, principal, review["id"], "new", virtual_operations=False,
                                    expected_version=review["version"], reason="Sin operaciones virtuales")
    assert item["virtualOperations"] is False
    publish(session, principal, 2, [invoice("Contribuyente como proveedor", "CON010101AA1", "MEX")])
    review = suppliers.list_reviews(session, principal.scope_id)[0]
    assert review["reason"] == "missing_identity"
    assert review["candidate"]["taxId"] == ""
    assert review["candidate"]["identityConflict"] == "contributor_rfc"
    assert all(s["taxId"] != "CON010101AA1" for s in suppliers.list_suppliers(session, principal.scope_id))


def test_new_505_supplier_uses_older_independent_virtual_evidence_and_tracks_its_retirement(domain):
    session, principal = domain
    first = publish(session, principal, 1, [], [
        {"clave_caso": "V1", "identificador_caso": "ABC010101AA1", "complemento_caso": "CN"}])
    assert suppliers.list_suppliers(session, principal.scope_id) == []  # 507 never creates a supplier.
    assert suppliers.list_reviews(session, principal.scope_id) == []
    publish(session, principal, 2, [invoice("Identificado por 505", "ABC010101AA1", "MEX")])
    supplier = suppliers.list_suppliers(session, principal.scope_id)[0]
    assert supplier["virtualOperations"] is True
    assert suppliers.list_reviews(session, principal.scope_id) == []
    detail = suppliers.get_supplier(session, principal.scope_id, supplier["id"])
    assert any(e["runId"] == first.id and e["virtualConfirmed"] for e in detail["evidence"])
    publish(session, principal, 1, [], [])
    assert suppliers.get_supplier(session, principal.scope_id, supplier["id"])["virtualOperations"] is None


def test_later_virtual_evidence_completes_pending_national_alta_without_manual_guess(domain):
    session, principal = domain
    publish(session, principal, 1, [invoice("Nacional identificado", "ABC010101AA1", "MEX")])
    assert suppliers.list_reviews(session, principal.scope_id)[0]["reason"] == "virtual_confirmation"
    publish(session, principal, 2, [], [
        {"clave_caso": "V1", "identificador_caso": "ABC010101AA1", "complemento_caso": "CN"}])
    assert suppliers.list_reviews(session, principal.scope_id) == []
    assert suppliers.list_suppliers(session, principal.scope_id)[0]["virtualOperations"] is True


def test_conflicting_data_explicit_new_and_correctable_equivalence(domain):
    session, principal = domain
    publish(session, principal, 1, [invoice()])
    existing = suppliers.list_suppliers(session, principal.scope_id)[0]
    publish(session, principal, 2, [invoice(tax="TAX200", calle_proveedor="500 Another St")])
    review = suppliers.list_reviews(session, principal.scope_id)[0]
    assert review["reason"] == "conflicting_data"
    created = suppliers.resolve_review(session, principal, review["id"], "new", expected_version=review["version"],
                                       reason="Sociedad distinta")
    assert created["id"] != existing["id"]
    alias = suppliers.list_equivalences(session, principal.scope_id)[0]
    suppliers.update_equivalence(session, principal, alias["id"], existing["id"], alias["version"],
                                  reason="Corregir identificación")
    assert any(e["period"] == "Febrero_2026" for e in suppliers.get_supplier(session, principal.scope_id, existing["id"])["evidence"])


@pytest.mark.parametrize("code,mode", [("V1", "CN"), ("V1", "RFE"), ("V1", "V2"), ("V2", ""), ("V8", ""), ("V5", ""), ("V6", "")])
def test_supported_507_rfc_modes_confirm_virtual(code, mode, domain):
    session, principal = domain
    publish(session, principal, 1, [invoice("Nacional", "ABC010101AA1", "MEX")], [
        {"clave_caso": code, "identificador_caso": "ABC010101AA1", "complemento_caso": mode}])
    assert suppliers.list_suppliers(session, principal.scope_id)[0]["virtualOperations"] is True


def test_report_immutable_pair_yellow_cycle_and_retry(domain):
    session, principal = domain
    publish(session, principal, 1, [invoice()])
    report = suppliers.create_report(session, principal, "report-month-001", "Primer reporte")
    before = suppliers.report_xlsx(session, principal.scope_id, report["id"], "foreign")
    book = load_workbook(BytesIO(before))
    assert book.active.max_column == 7
    assert [cell.value for cell in book.active[1]] == suppliers.DEFAULT_SCHEMAS["foreign"]
    assert book.active.cell(2, 7).value == "Alta"
    assert all(c.fill.fgColor.rgb == "FFFFFF00" for c in book.active[2])
    assert suppliers.create_report(session, principal, "report-month-001", "Primer reporte")["id"] == report["id"]
    assert len(suppliers.list_reports(session, principal.scope_id)) == 1
    second = suppliers.create_report(session, principal, "report-month-002", "Segundo reporte")
    later = load_workbook(BytesIO(suppliers.report_xlsx(session, principal.scope_id, second["id"], "foreign")))
    assert later.active.cell(2, 7).value == "Sin modificación"
    assert all(c.fill.fgColor.rgb != "FFFFFF00" for c in later.active[2])
    assert suppliers.report_xlsx(session, principal.scope_id, report["id"], "foreign") == before
    assert load_workbook(BytesIO(suppliers.report_xlsx(session, principal.scope_id, report["id"], "national"))).active.max_column == 3
    assert any(e.action == "suppliers.report_created" for e in session.scalars(select(AuditEvent)))


def test_pending_automatic_report_cancelled_when_any_captured_month_is_replaced(domain):
    session, principal = domain
    january = publish(session, principal, 1, [invoice()])
    suppliers.request_monthly_report(session, january, principal.id)
    initial = suppliers.monthly_report_state(session, principal.scope_id, january.id)
    assert initial["status"] == "ready"
    before = suppliers.report_xlsx(session, principal.scope_id, initial["report"]["id"], "foreign")
    february = publish(session, principal, 2, [invoice(name="ALAMO TECNOLOGIAS SA", tax="")])
    suppliers.request_monthly_report(session, february, principal.id)
    assert suppliers.monthly_report_state(session, principal.scope_id, february.id)["status"] == "pending_review"
    publish(session, principal, 1, [])
    assert suppliers.monthly_report_state(session, principal.scope_id, february.id)["status"] == "unavailable"
    request = next(r for r in suppliers.list_report_requests(session, principal.scope_id) if r["runId"] == february.id)
    assert request["status"] == "cancelled" and request["reason"] == "source_superseded"
    assert suppliers.report_xlsx(session, principal.scope_id, initial["report"]["id"], "foreign") == before


def test_national_review_completion_automatically_delivers_both_files(domain):
    session, principal = domain
    run = publish(session, principal, 1, [invoice("Proveedor nacional", "ABC010101AA1", "MEX")])
    suppliers.request_monthly_report(session, run, principal.id)
    assert suppliers.monthly_report_state(session, principal.scope_id, run.id)["status"] == "pending_review"
    review = suppliers.list_reviews(session, principal.scope_id)[0]
    suppliers.resolve_review(session, principal, review["id"], "new", virtual_operations=False,
                             expected_version=review["version"], reason="Alta nacional confirmada")
    state = suppliers.monthly_report_state(session, principal.scope_id, run.id)
    assert state["status"] == "ready"
    national = load_workbook(BytesIO(suppliers.report_xlsx(session, principal.scope_id, state["report"]["id"], "national")))
    assert [cell.value for cell in national.active[2]] == ["ABC010101AA1", "No", "Alta"]
    assert all(cell.fill.fgColor.rgb == "FFFFFF00" for cell in national.active[2])
    assert load_workbook(BytesIO(suppliers.report_xlsx(session, principal.scope_id, state["report"]["id"], "foreign"))).active.max_row == 1
