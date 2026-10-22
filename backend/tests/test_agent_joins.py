"""Complete related-table queries and Excel with real synthetic SQL data, no Azure."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import json
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.core.errors import ApplicationError
from app.modules.foundry import validate_tool_arguments
from app.modules.foundry.data_queries import describe_data, query_data
from app.modules.foundry.tools import tool_definitions
from app.persistence.business import BUSINESS_TABLES, HEADER_COLUMNS, normalize_header
from app.persistence.models import AgentArtifact, ProcessingFile, ProcessingTable, Run
from test_agent_reports import READER, request_chat, seed, specification, system as system


def declaration(number="0000001"):
    return {"Patente": "3658", "Pedimento": number, "SeccionAduanera": "730", "PedimentoCompleto": "25-73-3658-" + number}


def item(number="0000001", *, tariff="00123456", sequence="1", value="100"):
    return {**declaration(number), "Fraccion": tariff, "SecuenciaFraccion": sequence,
            "ValorDolares": value, "CantidadUMComercial": "3", "TipoOperacion": "1", "PaisOrigenDestino": "USA"}


def seed_tables(session, tables, **options):
    first = next(iter(tables))
    run_id = seed(session, first, tables[first], **options)
    run = session.get(Run, run_id)
    for ordinal, (code, rows) in enumerate(tables.items(), start=2):
        if code == first:
            continue
        file = ProcessingFile(run_id=run_id, period_id=run.period_id, file_name=f"{code}.asc", table_code=code,
                              ordinal=ordinal, sha256="0" * 64)
        session.add(file)
        session.flush()
        session.add(ProcessingTable(run_id=run_id, table_code=code, sheet_name=code, row_count=len(rows), headers_json="[]"))
        for index, row in enumerate(rows):
            values = {"period_id": run.period_id, "processing_run_id": run_id, "processing_file_id": file.id,
                      "archivo_origen": f"{code}.asc", "folio_origen": "0001", "table_code": code,
                      "source_row_number": index + 1, "original_dates_json": "{}"}
            values.update({HEADER_COLUMNS[code][normalize_header(header)]: value for header, value in row.items()})
            session.execute(BUSINESS_TABLES[code].insert().values(**values))
    session.flush()
    return run_id


def joined(**changes):
    return specification("551", columns=["551.Fraccion", "501.Rfc", "Valor MXN"],
                         joins=[{"tableCode": "501", "kind": "left"}],
                         calculations=[calculation()], **changes)


def calculation(left="551.ValorDolares", operator="multiply", right="501.TipoCambio", alias="Valor MXN", column=True):
    return {"left": left, "operator": operator, "right": right, "rightIsColumn": column, "alias": alias}


def plan(**changes):
    result = joined()
    result.update(changes)
    return result


def test_real_join_conversion_grouping_and_full_preview(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        run = seed_tables(session, {"551": [item(sequence=str(n), value="10") for n in range(125)],
                                    "501": [{**declaration(), "TipoCambio": "17.5", "Rfc": "TEST123"}]})
        detail = query_data(session, READER.scope_id, joined())
        assert len(detail.rows) == 125 and detail.rows[0] == ["00123456", "TEST123", Decimal("175")]
        assert detail.preview(1)["totalResultRows"] == 125 and detail.preview(1)["previewLimited"]
        assert detail.sources[0]["runId"] == run and detail.sources[0]["tables"] == ["551", "501"]
        assert detail.preview(1)["units"]["Valor MXN"] == "MXN"
        result = query_data(session, READER.scope_id, plan(columns=[], groupBy=["551.Fraccion"],
            metrics=[{"function": "sum", "column": "Valor MXN", "alias": "Total MXN"}],
            orderBy=[{"column": "Total MXN", "direction": "desc"}]))
        assert result.rows == [["00123456", Decimal("21875")]]
        assert result.scanned_rows == 126 and result.matched_rows == 125


def test_left_inner_missing_keys_and_unavailable_totals(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(), item("0000002")], "501": [{**declaration(), "TipoCambio": "17"}]})
        left = query_data(session, READER.scope_id, joined())
        assert {row[-1] for row in left.rows} == {Decimal("1700"), None}
        assert left.diagnostics["joins"][0]["unmatchedBaseRows"] == 1
        aggregate = query_data(session, READER.scope_id, plan(columns=[], metrics=[
            {"function": "sum", "column": "Valor MXN", "alias": "Total"}]))
        assert aggregate.rows == [[None]]  # no partial total omitting the unmatched item
        inner = query_data(session, READER.scope_id, plan(joins=[{"tableCode": "501", "kind": "inner"}]))
        assert inner.rows == [["00123456", None, Decimal("1700")]]


def test_join_uses_same_month_version_and_authenticated_scope(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": "17"}]})
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": "20"}]}, month=6)
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": "999"}]}, scope="foreign")
        # A retired version in the SAME period must not enter the related table index.
        active = session.scalar(select(Run).where(Run.scope_id == "authorized", Run.year == 2025))
        retired = Run(scope_id=active.scope_id, period_id=active.period_id, year=2025, kind="monthly",
                      input_fingerprint=uuid4().hex * 2, publication_status="published", status="completed")
        session.add(retired)
        session.flush()
        old_file = ProcessingFile(run_id=retired.id, period_id=active.period_id, file_name="old.asc",
                                  table_code="501", ordinal=1, sha256="0" * 64)
        session.add(old_file)
        session.flush()
        session.add(ProcessingTable(run_id=retired.id, table_code="501", sheet_name="501", row_count=1, headers_json="[]"))
        session.execute(BUSINESS_TABLES["501"].insert().values(
            period_id=active.period_id, processing_run_id=retired.id, processing_file_id=old_file.id,
            archivo_origen="old.asc", folio_origen="old", table_code="501", source_row_number=1,
            original_dates_json="{}", patente="3658", pedimento="0000001", seccion_aduanera="730", tipo_cambio="777"))
        session.flush()
        result = query_data(session, READER.scope_id, plan(periods=["Mayo_2025", "Junio_2025"],
            columns=["Periodo", "Valor MXN"]))
        assert dict(result.rows) == {"Mayo_2025": Decimal("1700"), "Junio_2025": Decimal("2000")}
        assert all(source["version"] == 2 for source in result.sources)


def test_qualified_base_columns_work_without_a_join_or_calculation(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(value="10"), item(sequence="2", value="20")]})
        result = query_data(session, READER.scope_id, plan(joins=[], calculations=[], columns=["551.ValorDolares"],
            orderBy=[{"column": "551.ValorDolares", "direction": "desc"}]))
        assert result.headers == ["551.ValorDolares"] and result.rows == [[Decimal("20")], [Decimal("10")]]


@pytest.mark.parametrize("bad", ["duplicate", "incomplete"])
def test_ambiguous_and_incomplete_target_keys_never_multiply_rows(system, bad):
    app, _ = system
    with app.state.session_factory.begin() as session:
        headers = [{**declaration(), "TipoCambio": "17"}] * 2 if bad == "duplicate" else [{"TipoCambio": "17"}]
        seed_tables(session, {"551": [item()], "501": headers})
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, joined())
        assert error.value.code == ("agent_join_ambiguous" if bad == "duplicate" else "agent_join_key_invalid")


def test_incomplete_base_key_is_not_silently_matched_or_discarded(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        incomplete = item()
        incomplete.pop("Patente")
        seed_tables(session, {"551": [incomplete], "501": [{**declaration(), "TipoCambio": "17"}]})
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, joined())
        assert error.value.code == "agent_join_key_invalid"


@pytest.mark.parametrize("function", ["sum", "mean"])
def test_coarser_measures_cannot_be_inflated_even_across_different_groups(system, function):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(), item(tariff="00999999", sequence="2")],
                              "501": [{**declaration(), "TipoCambio": "17", "TotalFletes": "100"}]})
        query = plan(columns=[], calculations=[], groupBy=["551.Fraccion"], metrics=[
            {"function": function, "column": "501.TotalFletes", "alias": "Flete"}])
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == "agent_join_repeated_measure"
        query["calculations"] = [calculation("501.TotalFletes", "add", "1", "Flete ajustado", False)]
        query["metrics"][0]["column"] = "Flete ajustado"
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == "agent_join_repeated_measure"


def test_payment_item_and_header_joins_preserve_tax_rows_and_verify_operation(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        tax = {**declaration(), "Fraccion": "00123456", "SecuenciaFraccion": "1", "FormaPago": "0", "ClaveContribucion": "6"}
        seed_tables(session, {"557": [{**tax, "ImportePago": "100"}, {**tax, "ImportePago": "200"}],
                              "551": [item()], "501": [{**declaration(), "TipoCambio": "20"}]})
        query = specification("557", columns=[], joins=[{"tableCode": "551", "kind": "inner"}, {"tableCode": "501", "kind": "left"}],
                              calculations=[calculation("557.ImportePago", "divide", "501.TipoCambio", "Pago USD")],
                              groupBy=["551.PaisOrigenDestino"],
                              filters=[{"column": "551.TipoOperacion", "operator": "eq", "value": "1"}],
                              metrics=[{"function": "sum", "column": "Pago USD", "alias": "Total USD"}])
        assert query_data(session, READER.scope_id, query).rows == [["USA", Decimal("15")]]
        query["calculations"] = []
        query["metrics"][0]["column"] = "551.ValorDolares"
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == "agent_join_repeated_measure"


@pytest.mark.parametrize("target", ["505", "557", "users", "551"])
def test_only_catalogued_many_to_one_relations_are_allowed(system, target):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()]})
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, plan(joins=[{"tableCode": target, "kind": "left"}]))
        assert error.value.code == "agent_join_forbidden"


def test_missing_join_table_and_combined_row_cap_are_not_partial_results(system, monkeypatch):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()]})
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, joined())
        assert error.value.code == "agent_table_unavailable"
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(), item(sequence="2")], "501": [{**declaration(), "TipoCambio": "17"}]}, month=6)
        monkeypatch.setattr("app.modules.foundry.data_queries.MAX_SOURCE_ROWS", 2)
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, plan(periods=["Junio_2025"]))
        assert error.value.code == "agent_query_too_large"


@pytest.mark.parametrize("operator,expected", [("add", "103"), ("subtract", "97"), ("multiply", "300"),
    ("divide", "33.3333333333333"), ("percent_of", "3333.33333333333"), ("percent_change", "3233.33333333333")])
def test_bounded_calculations_and_excel_numeric_precision(system, operator, expected):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()]})
        result = query_data(session, READER.scope_id, plan(joins=[], columns=["Calculated"],
            calculations=[calculation("ValorDolares", operator, "3", "Calculated", False)]))
        assert result.rows == [[Decimal(expected)]]


def test_zero_denominators_are_null_and_invalid_values_are_not_omitted(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()]})
        result = query_data(session, READER.scope_id, plan(joins=[], columns=["Ratio"],
            calculations=[calculation("ValorDolares", "divide", "0", "Ratio", False)]))
        assert result.rows == [[None]] and result.diagnostics["zeroDenominators"] == {"Ratio": 1}
        session.execute(BUSINESS_TABLES["551"].update().values(valor_dolares="bad"))
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, plan(joins=[], columns=["Ratio"],
                calculations=[calculation("ValorDolares", "divide", "2", "Ratio", False)]))
        assert error.value.code == "agent_invalid_source_value"


@pytest.mark.parametrize("rate", ["0", "-17"])
def test_currency_conversion_never_accepts_a_nonpositive_exchange_rate(system, rate):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": rate}]})
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, joined())
        assert error.value.code == "agent_invalid_source_value"


def test_currency_rules_apply_to_related_and_calculated_columns(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"505": [{**declaration(), "ValorMonedaExtranjera": "100", "MonedaFacturacion": "EUR", "ValorDolares": "120"}],
                              "501": [{**declaration(), "TotalFletes": "10"}]})
        query = specification("505", columns=[], joins=[{"tableCode": "501", "kind": "left"}],
            calculations=[calculation("ValorMonedaExtranjera", "multiply", "2", "Adjusted", False)],
            metrics=[{"function": "sum", "column": "Adjusted", "alias": "Total"}])
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == "agent_mixed_currencies"
        query["groupBy"] = ["MonedaFacturacion"]
        assert query_data(session, READER.scope_id, query).rows == [["EUR", Decimal("200")]]
        query["calculations"] = [calculation("ValorDolares", "add", "501.TotalFletes", "Mixed")]
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == "agent_mixed_units"


def test_calculations_can_be_filtered_and_chained_but_not_execute_code(system):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(value="10"), item(sequence="2", value="100")]})
        query = plan(joins=[], columns=["Tax"], calculations=[
            calculation("ValorDolares", "multiply", "0.16", "Tax", False),
            calculation("Tax", "add", "1", "Adjusted", False)],
            filters=[{"column": "Adjusted", "operator": "gt", "value": "5"}])
        assert query_data(session, READER.scope_id, query).rows == [[Decimal("16")]]
        for bad in ["__import__('os')", "=1+1", "NaN", "1e999999", "1,000"]:
            query["calculations"][0]["right"] = bad
            with pytest.raises(ApplicationError) as error:
                query_data(session, READER.scope_id, query)
            assert error.value.code == "agent_calculation_invalid"


def test_real_combined_excel_dates_numbers_units_sources_and_chat_evidence(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": "17",
            "FechaPagoReal": datetime(2025, 5, 15), "NombreContribuyente": "=HYPERLINK(\"https://bad.invalid\")"}]})
    query = plan(columns=["551.Fraccion", "501.FechaPagoReal", "501.NombreContribuyente", "Valor MXN"])
    args = {"title": "Combined", "sheets": [{"name": "Combined", "query": query}]}
    _, response, seen = request_chat(client, monkeypatch, "create_excel", args)
    assert response.status_code == 200
    file = response.json()["attachments"][0]
    book = load_workbook(BytesIO(client.get(file["downloadUrl"]).content))
    assert book["Combined"]["D2"].value == 1700 and book["Combined"]["D2"].data_type == "n"
    assert isinstance(book["Combined"]["B2"].value, datetime)
    assert book["Combined"]["C2"].data_type == "s"
    assert book["Fuentes"]["B2"].value == "551 + 501"
    assert seen["data"]["manifest"]["sheets"][0]["columnUnits"]["Valor MXN"] == "MXN"
    with app.state.session_factory() as session:
        manifest = json.loads(session.scalar(select(AgentArtifact)).manifest_json)
        assert manifest["sheets"][0]["diagnostics"]["joins"][0]["matchedBaseRows"] == 1
    _, response, seen = request_chat(client, monkeypatch, "query_data", {"query": query, "limit": 1})
    assert response.status_code == 200
    assert seen["evidence"][0]["diagnostics"]["joins"][0]["unmatchedBaseRows"] == 0


def test_combined_workbook_exports_all_detail_and_unions_sources_across_sheets(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item(sequence=str(n), value="10") for n in range(125)],
            "501": [{**declaration(), "TipoCambio": "17.5", "Rfc": "TEST123"}],
            "505": [{**declaration(), "ProveedorMercancia": "Vendor", "ValorDolares": "1250"}]})
    args = {"title": "Complete", "sheets": [
        {"name": "Detail", "query": joined()},
        {"name": "Invoices", "query": specification()},
        {"name": "Summary", "query": plan(columns=[], groupBy=["551.Fraccion"],
            metrics=[{"function": "sum", "column": "Valor MXN", "alias": "Total MXN"}])}]}
    _, response, seen = request_chat(client, monkeypatch, "create_excel", args)
    assert response.status_code == 200
    book = load_workbook(BytesIO(client.get(response.json()["attachments"][0]["downloadUrl"]).content))
    assert book["Detail"].max_row == 126 and book["Detail"]["C126"].value == 175
    assert book["Summary"]["B2"].value == 21875 and book["Summary"]["B2"].data_type == "n"
    assert seen["data"]["manifest"]["sources"][0]["tables"] == ["501", "505", "551"]


@pytest.mark.parametrize("invalid", ["column", "alias", "reference"])
def test_related_columns_and_aliases_are_checked_against_the_closed_catalog(system, invalid):
    app, _ = system
    with app.state.session_factory.begin() as session:
        seed_tables(session, {"551": [item()], "501": [{**declaration(), "TipoCambio": "17"}]})
        query = joined()
        if invalid == "column":
            query["columns"] = ["501.Password"]
        elif invalid == "alias":
            query["calculations"][0]["alias"] = "ValorDolares"
        else:
            query["calculations"][0]["right"] = "551.ValorDolares; DROP TABLE users"
        with pytest.raises(ApplicationError) as error:
            query_data(session, READER.scope_id, query)
        assert error.value.code == ("agent_column_duplicate" if invalid == "alias" else "agent_column_unknown")


def test_catalog_and_registered_schemas_expose_relations_without_identity_or_sql():
    assert describe_data("551")["relationships"][0]["tableCode"] == "501"
    assert {relation["tableCode"] for relation in describe_data("557")["relationships"]} == {"501", "551"}
    assert all(relation["tableCode"] != "551" for relation in describe_data("505")["relationships"])
    for tool in tool_definitions(False):
        for schema in [tool["parameters"], *tool["parameters"].get("$defs", {}).values()]:
            if schema.get("type") == "object":
                assert set(schema["required"]) == set(schema["properties"])
                assert schema["additionalProperties"] is False
    invalid = plan()
    invalid["joins"][0]["on"] = "1=1"
    with pytest.raises(ApplicationError):
        validate_tool_arguments("query_data", {"query": invalid, "limit": 1})
