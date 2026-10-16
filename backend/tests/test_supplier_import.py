"""Both selected workbooks are imported atomically into the caller's scope."""
from io import BytesIO

import xlsxwriter

from app.modules.identity.auth import Principal, current_principal
from app.modules.suppliers.service import DEFAULT_SCHEMAS
from test_api_worker import system as system


def workbook(kind, rows):
    content = BytesIO()
    with xlsxwriter.Workbook(content, {"in_memory": True, "strings_to_formulas": False}) as book:
        sheet = book.add_worksheet("Sheet1")
        sheet.write_row(0, 0, DEFAULT_SCHEMAS[kind])
        for index, row in enumerate(rows, 1):
            sheet.write_row(index, 0, row)
    return content.getvalue()


def files(foreign=None, national=None):
    return {
        "foreign": ("foreign.xlsx", foreign or workbook("foreign", [["Precision Components Limited", "", "", "Estados Unidos de América", "Main St", "Proveedor", "Alta"]])),
        "national": ("national.xlsx", national or workbook("national", [["PVE010101AB1", "Si", "Sin modificación"]])),
    }


def test_selected_xlsx_pair_merges_existing_catalog_and_is_retry_safe(system):
    _, client, _ = system
    created = client.post("/api/v1/suppliers", json={"kind": "foreign", "name": "Precision Components LTD", "country": "USA", "taxId": "US112233"})
    assert created.status_code == 201
    result = client.post("/api/v1/suppliers/import", files=files())
    assert result.status_code == 200, result.text
    assert result.json()["initialized"] is True
    assert result.json()["national"] == result.json()["foreign"] == 1
    assert client.post("/api/v1/suppliers/import", files=files()).json()["catalog"] == 2
    catalog = client.get("/api/v1/suppliers?type=foreign").json()
    assert catalog["items"][0]["id"] == created.json()["id"]
    detail = client.get(f"/api/v1/suppliers/{created.json()['id']}").json()
    seed = next(e for e in detail["evidence"] if e["sourceKind"] == "seed")
    assert seed["records"][0]["rowNumber"] == 2
    assert seed["records"][0]["values"]["Nombre o Razón social"] == "Precision Components Limited"


def test_invalid_second_file_never_leaves_a_partial_catalog(system):
    _, client, _ = system
    invalid = workbook("national", [["PVE010101AB1", "Maybe", "Alta"]])
    result = client.post("/api/v1/suppliers/import", files=files(national=invalid))
    assert result.status_code == 400, result.text
    assert client.get("/api/v1/suppliers/summary").json()["catalog"] == 0
    assert client.get("/api/v1/suppliers/summary").json()["initialized"] is False


def test_initial_rows_with_similar_legal_suffixes_are_preserved_verbatim(system):
    _, client, _ = system
    foreign = workbook("foreign", [
        ["Precision Components Limited", "", "", "Estados Unidos de América", "Branch A", "Proveedor", "Alta"],
        ["Precision Components LTD", "", "", "Estados Unidos de América", "Branch B", "Proveedor", "Sin modificación"],
    ])
    result = client.post("/api/v1/suppliers/import", files=files(foreign=foreign))
    assert result.status_code == 200, result.text
    assert result.json()["foreign"] == 2
    assert {s["name"] for s in client.get("/api/v1/suppliers?type=foreign").json()["items"]} == {
        "Precision Components Limited", "Precision Components LTD"}


def test_empty_national_template_is_valid_and_does_not_block_foreign_workbook(system):
    _, client, _ = system
    response = client.post("/api/v1/suppliers/import", files=files(national=workbook("national", [])))
    assert response.status_code == 200, response.text
    assert response.json()["national"] == 0 and response.json()["foreign"] == 1


def test_imported_format_and_auxiliary_sheet_survive_automatic_exports_and_retries(system):
    from openpyxl import load_workbook
    from test_supplier_monthly_automatic import publish, state
    _, client, _ = system
    content = BytesIO()
    with xlsxwriter.Workbook(content, {"in_memory": True}) as book:
        sheet = book.add_worksheet("Original suppliers")
        aux = book.add_worksheet("Reference")
        aux.write(0, 0, "Original reference")
        header = book.add_format({"font_name": "Calibri", "font_size": 12, "bold": True, "bg_color": "#808080"})
        body = book.add_format({"font_name": "Calibri", "font_size": 11, "bg_color": "#FFFF00"})
        sheet.write_row(0, 0, DEFAULT_SCHEMAS["foreign"], header)
        sheet.write_row(1, 0, ["Precision Components Limited", "", "", "", "", "Proveedor", "Sin modificación"], body)
        sheet.set_column(0, 0, 58)
        sheet.set_row(0, 27)
        sheet.set_row(1, 23)
        sheet.freeze_panes(1, 0)
        sheet.data_validation("F2:F500", {"validate": "list", "source": ["Proveedor", "Cliente"]})
    result = client.post("/api/v1/suppliers/import", files=files(foreign=content.getvalue()))
    assert result.status_code == 200, result.text
    run_id = publish(system, 1, ["A New Supplier Limited"])
    report_id = state(client, run_id)["report"]["id"]
    payload = client.get(f"/api/v1/suppliers/reports/{report_id}/foreign.xlsx").content
    book = load_workbook(BytesIO(payload))
    sheet = book.worksheets[0]
    assert book.sheetnames == ["Original suppliers", "Reference"]
    assert book["Reference"]["A1"].value == "Original reference"
    assert sheet["A1"].font.sz == 12 and sheet["A1"].fill.fgColor.rgb == "FF808080"
    assert sheet.row_dimensions[1].height == 27
    assert sheet.row_dimensions[3].height == 23
    assert sheet.column_dimensions["A"].width > 58
    assert sheet.freeze_panes == "A2" and len(sheet.data_validations.dataValidation) == 1
    assert sheet["A2"].value == "Precision Components Limited"
    assert sheet["A2"].fill.fgColor.rgb != "FFFFFF00"
    assert sheet["A3"].value == "A New Supplier Limited"
    assert all(c.fill.fgColor.rgb == "FFFFFF00" for c in sheet[3])
    # Changing the current catalog must never change already published bytes.
    client.post("/api/v1/suppliers/sync", json={})
    assert client.get(f"/api/v1/suppliers/reports/{report_id}/foreign.xlsx").content == payload


def test_reader_and_other_scope_cannot_initialize_someone_elses_catalog(system):
    app, client, _ = system
    app.dependency_overrides[current_principal] = lambda: Principal("reader", "Reader", frozenset({"Reader"}), "local")
    assert client.post("/api/v1/suppliers/import", files=files()).status_code == 403
    app.dependency_overrides[current_principal] = lambda: Principal("admin", "Admin", frozenset({"Admin"}), "other-scope")
    assert client.post("/api/v1/suppliers/import", files=files()).status_code == 200
    app.dependency_overrides.clear()
    assert client.get("/api/v1/suppliers/summary").json()["catalog"] == 0


def test_invalid_xls_returns_a_useful_error_and_does_not_change_catalog(system):
    _, client, _ = system
    uploaded = files()
    uploaded["foreign"] = ("foreign.xls", b"not an Excel file")
    result = client.post("/api/v1/suppliers/import", files=uploaded)
    assert result.status_code == 400, result.text
    assert result.json()["title"] == "INVALID_SUPPLIER_WORKBOOK"
    assert client.get("/api/v1/suppliers/summary").json()["catalog"] == 0
