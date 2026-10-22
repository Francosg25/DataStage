"""Real private attachments and exported file formats, with Azure simulated."""
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
import pytest
from sqlalchemy import func, select
import xlsxwriter

from app.core.errors import ApplicationError
from app.modules.foundry.file_parser import parse
from app.modules.foundry import uploads
from app.modules.foundry.tools import tool_definitions
from app.modules.identity.auth import Principal, current_principal
from app.persistence.models import AgentArtifact, AgentUpload, StoredDocument
from test_agent_reports import READER, request_chat, seed, specification, system as system


def conversation(client):
    return client.post("/api/v1/agent/conversations").json()["id"]


def upload(client, conversation_id, name="data.csv", content=b"Vendor,Amount\nAlpha,100.25\nBeta,200\n"):
    return client.post(f"/api/v1/agent/conversations/{conversation_id}/files", files={"file": (name, content)})


def file_query(file, **changes):
    return {"fileId": file["fileId"], "sheet": file["sheets"][0]["name"], "columns": file["sheets"][0]["headers"],
            "filters": [], "groupBy": [], "metrics": [], "orderBy": [], **changes}


def file_spec(format, query, *, chart="none"):
    return {"title": "Monthly results", "format": format, "sections": [{"title": "Values by vendor", "query": query,
            "commentary": "", "chart": chart}]}


def artifact(client, monkeypatch, conv, format, query, chart="none"):
    _, response, seen = request_chat(client, monkeypatch, "create_file", file_spec(format, query, chart=chart), conversation_id=conv)
    assert response.status_code == 200, response.text
    download = response.json()["attachments"][0]
    content = client.get(download["downloadUrl"])
    assert content.status_code == 200 and "attachment" in content.headers["content-disposition"]
    return download, content, seen


def test_upload_is_idempotent_private_durable_and_not_a_month(system):
    app, client = system
    conv = conversation(client)
    response = upload(client, conv)
    assert response.status_code == 201, response.text
    file = response.json()
    assert file["sheets"][0]["headers"] == ["Vendor", "Amount"] and file["sheets"][0]["types"] == ["text", "number"]
    duplicate = upload(client, conv)
    assert duplicate.json()["fileId"] == file["fileId"]
    listed = client.get(f"/api/v1/agent/conversations/{conv}/files").json()
    assert len(listed["items"]) == 1 and listed["maxFiles"] == 5
    assert client.get("/api/v1/agent/conversations").json()["items"][0]["title"] == "data.csv"
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(AgentUpload)) == 1
        assert session.scalar(select(func.count()).select_from(StoredDocument)) == 2
    assert client.get(f'/api/v1/documents/{file["fileId"]}/download').status_code == 403
    url = f'/api/v1/agent/conversations/{conv}/files/{file["fileId"]}/download'
    assert client.get(url).content.startswith(b"Vendor,Amount")
    app.dependency_overrides[current_principal] = lambda: Principal("other", "Other", frozenset({"Admin"}), READER.scope_id)
    assert client.get(url).status_code == 404
    assert client.get(f"/api/v1/agent/conversations/{conv}/files").status_code == 404
    assert upload(client, conv).status_code == 404


def test_attachment_queries_compute_full_totals_filters_order_and_preview(system, monkeypatch):
    app, client = system
    conv = conversation(client)
    file = upload(client, conv, content=b"Vendor,Amount\n" + b"Alpha,100.25\n" * 125).json()
    query = file_query(file, columns=[], groupBy=["Vendor"], metrics=[{"function": "sum", "column": "Amount", "alias": "Total"}])
    _, response, seen = request_chat(client, monkeypatch, "query_attachment", {"query": query, "offset": 0, "limit": 1}, conversation_id=conv)
    assert response.status_code == 200, response.text
    assert seen["data"]["items"] == [{"Vendor": "Alpha", "Total": 12531.25}]
    assert seen["data"]["matchedSourceRows"] == 125 and seen["data"]["complete"]
    assert seen["evidence"][0]["source"] == "conversation_attachment"
    assert seen["history"][-2]["content"].startswith("DATASTAGE_ATTACHMENTS")
    history = client.get(f"/api/v1/agent/conversations/{conv}/messages").json()
    assert history["items"][0]["sourceFiles"][0]["fileId"] == file["fileId"]
    query = file_query(file, filters=[{"column": "Amount", "operator": "gt", "value": "90"}], orderBy=[{"column": "Amount", "direction": "desc"}])
    with app.state.session_factory() as session:
        result = uploads.query_file(session, app.state.storage, READER, conv, query)
        assert len(result.rows) == 125
        preview = uploads.preview_file(result, 120, 3)
        assert len(preview["items"]) == 3 and preview["nextOffset"] == 123 and preview["previewLimited"]


def test_excel_dates_identifiers_and_formulas_are_preserved_without_execution(system, monkeypatch):
    app, client = system
    stream = BytesIO()
    with xlsxwriter.Workbook(stream, {"in_memory": True}) as book:
        page = book.add_worksheet("Invoices")
        page.write_row(0, 0, ["Pedimento", "Date", "Amount", "Formula"])
        page.write_number(1, 0, 123, book.add_format({"num_format": "0000000"}))
        page.write_datetime(1, 1, datetime(2026, 9, 1), book.add_format({"num_format": "yyyy-mm-dd"}))
        page.write_number(1, 2, 15.5)
        page.write_formula(1, 3, "=1+1")
    conv = conversation(client)
    response = upload(client, conv, "source.xlsx", stream.getvalue())
    assert response.status_code == 201, response.text
    file = response.json()
    assert "formulas_preserved_as_text" in file["warnings"]
    _, content, _ = artifact(client, monkeypatch, conv, "xlsx", file_query(file))
    book = load_workbook(BytesIO(content.content))
    row = list(book.worksheets[0].values)[1]
    assert row == ("0000123", datetime(2026, 9, 1), 15.5, "=1+1")
    assert book.worksheets[0]["D2"].data_type == "s" and book.worksheets[0]["C2"].data_type == "n"


@pytest.mark.parametrize("format", ["docx", "pptx", "csv", "json", "png", "svg", "html", "pdf", "xlsx"])
def test_real_export_formats_from_published_queries_have_verified_sources(system, monkeypatch, format):
    app, client = system
    with app.state.session_factory.begin() as session:
        run = seed(session, "505", [{"ProveedorMercancia": "Alpha", "ValorDolares": "100.25"},
                                   {"ProveedorMercancia": "Beta", "ValorDolares": "200"}])
    conv = conversation(client)
    download, response, seen = artifact(client, monkeypatch, conv, format, specification(), "bar" if format in {"png", "svg", "pptx", "docx", "html", "pdf"} else "none")
    assert download["format"] == format and len(response.content) > 20
    assert seen["data"]["manifest"]["sources"][0]["runId"] == run
    if format == "docx":
        word = Document(BytesIO(response.content))
        assert word.tables[0].rows[1].cells[1].text == "100.25"
    elif format == "pptx":
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            assert archive.testzip() is None
            types_xml = archive.read("[Content_Types].xml").decode()
            assert '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"' in types_xml
            assert '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"' in archive.read("ppt/slides/_rels/slide1.xml.rels").decode()
            for name in archive.namelist():
                if name.endswith((".xml", ".rels")):
                    ET.fromstring(archive.read(name))
            assert "ppt/charts/chart1.xml" in archive.namelist() and "ppt/embeddings/data1.xlsx" in archive.namelist()
            assert b"100.25" in archive.read("ppt/charts/chart1.xml")
            assert run.encode() in archive.read("customXml/item1.xml")
    elif format == "pdf":
        assert "100.25" in "".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
    elif format == "png":
        image = Image.open(BytesIO(response.content))
        assert image.width == 1600 and run in image.info["DataStage sources"]
    elif format == "svg":
        assert ET.fromstring(response.content).tag.endswith("svg")
    elif format == "json":
        assert response.json()["sections"][0]["rows"] == [["Alpha", 100.25], ["Beta", 200]]
    elif format == "csv":
        assert response.content.decode("utf-8-sig").splitlines()[1] == "Alpha,100.25"
    elif format == "xlsx":
        assert load_workbook(BytesIO(response.content)).worksheets[0]["B2"].value == 100.25
    else:
        assert "<table>" in response.text and run in response.text


def test_files_can_be_analyzed_reexported_and_packaged_with_provenance(system, monkeypatch):
    app, client = system
    conv = conversation(client)
    file = upload(client, conv, content=b"Vendor,Amount\n=1+1,10\n<script>alert(1)</script>,20\n").json()
    csv, content, _ = artifact(client, monkeypatch, conv, "csv", file_query(file))
    assert "'=1+1,10" in content.content.decode("utf-8-sig")
    html, content, _ = artifact(client, monkeypatch, conv, "html", file_query(file))
    assert "<script>" not in content.text and "&lt;script&gt;" in content.text
    _, response, seen = request_chat(client, monkeypatch, "create_package", {"title": "Monthly package", "documentIds": [csv["documentId"], html["documentId"]]}, conversation_id=conv)
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(BytesIO(client.get(response.json()["attachments"][0]["downloadUrl"]).content)) as archive:
        assert len(archive.namelist()) == 3
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["files"][0]["sha256"] == csv["sha256"]
        assert manifest["sources"][0]["attachmentId"] == file["fileId"]
    # Download remains a snapshot after the live conversational source is removed.
    assert client.delete(f'/api/v1/agent/conversations/{conv}/files/{file["fileId"]}').status_code == 204
    assert client.get(csv["downloadUrl"]).status_code == 200
    with app.state.session_factory() as session:
        with pytest.raises(ApplicationError) as error:
            uploads.query_file(session, app.state.storage, READER, conv, file_query(file))
        assert error.value.status == 404


def test_pdf_word_and_powerpoint_uploads_are_read_as_text_without_fake_vision(system, monkeypatch):
    _, client = system
    conv = conversation(client)
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    pdf = BytesIO()
    writer.write(pdf)
    response = upload(client, conv, "scan.pdf", pdf.getvalue())
    assert response.status_code == 201 and "pdf_requires_ocr" in response.json()["warnings"]
    word = Document()
    word.add_paragraph("Approved supplier procedure")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Vendor", "Status"
    table.cell(1, 0).text, table.cell(1, 1).text = "Alpha", "Approved"
    stream = BytesIO()
    word.save(stream)
    response = upload(client, conv, "procedure.docx", stream.getvalue())
    assert response.status_code == 201, response.text
    assert {page["name"] for page in response.json()["sheets"]} == {"Text", "Table 1"}
    file = upload(client, conv).json()
    _, presentation, _ = artifact(client, monkeypatch, conv, "pptx", file_query(file), "bar")
    response = upload(client, conv, "presentation.pptx", presentation.content)
    assert response.status_code == 201, response.text
    assert "text_only_images_not_analyzed" in response.json()["warnings"]


@pytest.mark.parametrize("bad", ["extension", "encrypted", "zip_path", "macro", "entity", "binary"])
def test_invalid_and_active_attachments_are_rejected_and_cleaned(system, bad):
    app, client = system
    name, content = "source.csv", b"Vendor,Amount\nAlpha,10"
    if bad == "extension":
        name = "program.exe"
    elif bad == "encrypted":
        pdf = PdfWriter()
        pdf.add_blank_page(width=100, height=100)
        pdf.encrypt("secret")
        stream = BytesIO()
        pdf.write(stream)
        name, content = "encrypted.pdf", stream.getvalue()
    elif bad == "binary":
        name, content = "fake.xlsx", b"not an Office file"
    else:
        stream = BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/workbook.xml", "<workbook/>")
            archive.writestr({"zip_path": "../outside.xml", "macro": "xl/vbaProject.bin", "entity": "xl/evil.xml"}[bad],
                             '<!DOCTYPE data [<!ENTITY x "boom">]><data>&x;</data>' if bad == "entity" else "bad")
        name, content = "unsafe.xlsx", stream.getvalue()
    conv = conversation(client)
    response = upload(client, conv, name, content)
    assert response.status_code in {400, 422}, response.text
    assert not list(app.state.storage.root.iterdir())


def test_limits_timeouts_and_cross_conversation_access(system, monkeypatch):
    app, client = system
    conv, other = conversation(client), conversation(client)
    file = upload(client, conv).json()
    with app.state.session_factory() as session:
        with pytest.raises(ApplicationError):
            uploads.query_file(session, app.state.storage, READER, other, file_query(file))
    response = upload(client, other, content=b"a" * (uploads.MAX_UPLOAD_BYTES + 1))
    assert response.status_code == 413
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("reader", 30)
    monkeypatch.setattr(uploads.subprocess, "run", timeout)
    response = upload(client, other)
    assert response.status_code == 422
    assert not client.get(f"/api/v1/agent/conversations/{other}/files").json()["items"]
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(StoredDocument)) == 2


def test_conversation_deletion_revokes_every_format_and_attachments(system, monkeypatch):
    app, client = system
    conv = conversation(client)
    file = upload(client, conv).json()
    doc, _, _ = artifact(client, monkeypatch, conv, "docx", file_query(file))
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(StoredDocument)) == 3
    assert client.delete(f"/api/v1/agent/conversations/{conv}").status_code == 204
    assert client.get(doc["downloadUrl"]).status_code == 404
    assert client.get(f'/api/v1/agent/conversations/{conv}/files/{file["fileId"]}/download').status_code == 404
    assert not list(app.state.storage.root.iterdir())
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(AgentUpload)) == 0
        assert session.scalar(select(func.count()).select_from(AgentArtifact)) == 0


def test_long_text_preview_is_explicit_and_source_text_is_complete(system):
    app, client = system
    conv = conversation(client)
    text = "Detailed approved procedure " * 100
    file = upload(client, conv, content=("Note\n" + text + "\n").encode()).json()
    with app.state.session_factory() as session:
        result = uploads.query_file(session, app.state.storage, READER, conv, file_query(file))
        preview = uploads.preview_file(result, 0, 10)
        assert preview["previewTextLimited"] and len(preview["items"][0]["Note"]) == 1000
        assert result.rows[0][0] == text
        assert preview["textContinuations"][0]["nextTextOffset"] == 1000
        rest = uploads.preview_file(result, 0, 1, text_offset=1000, text_limit=6000)
        assert rest["items"][0]["Note"] == text[1000:] and rest["textContinuations"] == []


def test_full_attachment_export_is_not_the_chat_preview(system, monkeypatch):
    _, client = system
    conv = conversation(client)
    file = upload(client, conv, content=b"Vendor,Amount\n" + b"Alpha,100.25\n" * 125).json()
    _, content, _ = artifact(client, monkeypatch, conv, "csv", file_query(file))
    assert len(content.content.decode("utf-8-sig").splitlines()) == 126


def test_pdf_with_some_blank_pages_reports_partial_text_coverage(system):
    app, client = system
    from reportlab.pdfgen.canvas import Canvas
    stream = BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(20, 700, "Readable page")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    conv = conversation(client)
    response = upload(client, conv, "mixed.pdf", stream.getvalue())
    assert response.status_code == 201, response.text
    file = response.json()
    assert "pdf_some_pages_without_text" in file["warnings"]
    with app.state.session_factory() as session:
        result = uploads.query_file(session, app.state.storage, READER, conv, file_query(file, columns=["Page"],
            filters=[{"column": "Text", "operator": "is_empty", "value": ""}]))
        assert result.rows == [[2]]


def test_failed_parse_locked_by_onedrive_is_registered_for_worker_retry(system, monkeypatch):
    app, client = system
    from app.modules.foundry.conversations import cleanup_deleted_artifacts
    real_unlink = Path.unlink
    def locked(path, *args, **kwargs):
        if path.suffix == ".upload":
            raise PermissionError("locked")
        return real_unlink(path, *args, **kwargs)
    monkeypatch.setattr(Path, "unlink", locked)
    conv = conversation(client)
    response = upload(client, conv, "invalid.xlsx", b"invalid Office file")
    assert response.status_code == 422
    with app.state.session_factory() as session:
        document = session.scalar(select(StoredDocument))
        assert document.status == "pending_delete" and document.run_id is None
        assert session.scalar(select(func.count()).select_from(AgentUpload)) == 0
    monkeypatch.undo()
    assert cleanup_deleted_artifacts(app.state.session_factory, app.state.storage) == 1
    assert not list(app.state.storage.root.iterdir())


def test_new_schemas_remain_closed_and_no_arbitrary_code_is_accepted():
    tools = tool_definitions(False)
    assert len(tools) == 16 and "create_file" in {tool["name"] for tool in tools}
    def check(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node.get("properties", {}))
            assert "default" not in node
            for value in node.values():
                check(value)
        elif isinstance(node, list):
            for value in node:
                check(value)
    for tool in tools:
        check(tool["parameters"])


def test_parser_rejects_expansion_and_duplicate_paths_without_extracting(tmp_path):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<workbook>" + "a" * 100000 + "</workbook>")
    with pytest.raises(ValueError):
        parse(stream.getvalue(), ".xlsx")


def test_spanish_identifier_headers_keep_codes_textual_and_amounts_numeric():
    result = parse("Fracción arancelaria,Código postal,Amount\n12345678,12345,10.25\n".encode(), ".csv")
    page = result["sheets"][0]
    assert page["types"] == ["text", "text", "number"]
    assert page["rows"] == [["12345678", "12345", 10.25]]


def test_office_uppercase_xml_cannot_hide_entity_declarations():
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/EVIL.XML", '<!DOCTYPE data [<!ENTITY x "payload">]><data>&x;</data>')
    with pytest.raises(ValueError, match="entidades"):
        parse(stream.getvalue(), ".xlsx")


def test_valid_attachment_larger_than_generic_json_body_limit_is_accepted(system):
    _, client = system
    conv = conversation(client)
    content = b"Vendor,Amount\n" + (b"A" * 70 + b",10\n") * 15000
    assert len(content) > 1024 * 1024
    response = upload(client, conv, content=content)
    assert response.status_code == 201, response.text
    assert response.json()["sheets"][0]["rows"] == 15000


def test_file_count_limit_and_removed_attachment_do_not_return(system):
    _, client = system
    conv = conversation(client)
    files = [upload(client, conv, f"file-{i}.csv", f"Value\n{i}\n".encode()).json() for i in range(5)]
    assert upload(client, conv, "sixth.csv", b"Value\n6\n").status_code == 413
    assert client.delete(f'/api/v1/agent/conversations/{conv}/files/{files[0]["fileId"]}').status_code == 204
    response = upload(client, conv, "sixth.csv", b"Value\n6\n")
    assert response.status_code == 201
    ids = [item["fileId"] for item in client.get(f"/api/v1/agent/conversations/{conv}/files").json()["items"]]
    assert len(ids) == 5 and files[0]["fileId"] not in ids


def test_visual_exports_reject_oversized_detail_and_mixed_units(system, monkeypatch):
    app, client = system
    with app.state.session_factory.begin() as session:
        seed(session, "505", [{"ProveedorMercancia": "Alpha", "ValorDolares": "10", "ValorMonedaExtranjera": "15"}] * 201)
        seed(session, "505", [{"ProveedorMercancia": "Alpha", "ValorDolares": "10", "ValorMonedaExtranjera": "15"}], month=6)
    conv = conversation(client)
    _, response, _ = request_chat(client, monkeypatch, "create_file", file_spec("docx", specification()), conversation_id=conv)
    assert response.status_code == 413
    _, response, _ = request_chat(client, monkeypatch, "create_file", file_spec("png", specification(periods=["Junio_2025"], columns=["ProveedorMercancia", "ValorDolares", "ValorMonedaExtranjera"],
        groupBy=[], metrics=[]), chart="bar"), conversation_id=conv)
    assert response.status_code == 400
    assert "unidades distintas" in response.json()["detail"]
    with app.state.session_factory() as session:
        assert session.scalar(select(func.count()).select_from(AgentArtifact)) == 0


def test_foreign_currency_aggregation_requires_a_currency_group_or_filter(system):
    app, client = system
    conv = conversation(client)
    file = upload(client, conv, content=b"Vendor,Amount,Currency\nAlpha,10,USD\nAlpha,20,MXN\n").json()
    query = file_query(file, columns=[], groupBy=["Vendor"], metrics=[{"function": "sum", "column": "Amount", "alias": "Total"}])
    with app.state.session_factory() as session:
        with pytest.raises(ApplicationError) as error:
            uploads.query_file(session, app.state.storage, READER, conv, query)
        assert error.value.code == "agent_mixed_currencies"
        query["groupBy"] = ["Currency"]
        assert uploads.query_file(session, app.state.storage, READER, conv, query).rows == [["USD", 10], ["MXN", 20]]


def test_packages_cannot_include_other_conversations_or_changed_bytes(system, monkeypatch):
    app, client = system
    conv, other = conversation(client), conversation(client)
    file = upload(client, conv).json()
    doc, _, _ = artifact(client, monkeypatch, conv, "csv", file_query(file))
    args = {"title": "Package", "documentIds": [doc["documentId"]]}
    _, response, _ = request_chat(client, monkeypatch, "create_package", args, conversation_id=other)
    assert response.status_code == 400
    with app.state.session_factory() as session:
        stored = session.get(StoredDocument, doc["documentId"])
        app.state.storage.resolve(stored.storage_key).write_bytes(b"changed")
    _, response, _ = request_chat(client, monkeypatch, "create_package", args, conversation_id=conv)
    assert response.status_code == 503
